"""Audit saved products and provenance without refitting or modifying inputs."""

import ast
import hashlib
import json
import re
import subprocess
from collections import Counter
from pathlib import Path
from uuid import uuid4

import numpy as np
import pandas as pd
from obspy import UTCDateTime
from openpyxl import load_workbook


def sha(p):
    h = hashlib.sha256()
    with Path(p).open("rb") as f:
        for b in iter(lambda: f.read(1024 * 1024), b""):
            h.update(b)
    return h.hexdigest()


def main():
    ROOT = Path.cwd()
    OUT = ROOT / "results/VALIDATION/PROJECT_AUDIT" / ("run_" + uuid4().hex[:12])
    OUT.mkdir(parents=True, exist_ok=False)
    issues = []

    def record(condition, category, detail):
        if not condition:
            issues.append(dict(category=category, detail=detail))

    tracked = [
        Path(p)
        for p in subprocess.check_output(["git", "ls-files"], text=True).splitlines()
    ]
    production = [
        p
        for root in [
            ROOT / "results/NETWORKS",
            ROOT / "results/TEST",
            ROOT / "exports",
            Path.home() / "Desktop/Data_Bruit_FM/Data/YV",
            Path.home() / "Desktop/Data_Bruit_FM/Data/YL",
        ]
        for p in root.rglob("*")
        if p.is_file()
    ]
    before = {str(p): (p.stat().st_size, p.stat().st_mtime_ns) for p in production}
    print("Baseline", len(before), "production/source files", flush=True)
    master = pd.read_csv(
        "exports/master_inventory.csv", dtype=str, keep_default_na=False
    )
    protected = {
        p: sha(p)
        for p in [
            Path("exports/master_inventory.csv"),
            Path("exports/master_inventory.xlsx"),
        ]
    }
    for p in set(master.inventory_file):
        record(Path(p).is_file(), "missing_inventory", p)
        if Path(p).is_file():
            protected[Path(p)] = sha(p)
    parse = Counter()
    broken = []
    for p in tracked:
        if p.suffix == ".py":
            ast.parse(p.read_text())
            parse["python"] += 1
        elif p.suffix == ".json":
            json.loads(p.read_text())
            parse["json"] += 1
        elif p.suffix == ".md":
            body = p.read_text()
            parse["markdown"] += 1
            for ref in re.findall(r"\]\(([^)\s]+)\)", body):
                if re.match(r"^(https?://|mailto:|#)", ref):
                    continue
                target = ref.split("#")[0]
                if target and not (p.parent / target).exists():
                    broken.append(dict(file=str(p), target=ref))
    print(
        "Tracked file parsing",
        dict(parse),
        "broken local links",
        len(broken),
        flush=True,
    )
    workflows = []
    ray_paths = set()
    artifacts = 0
    for p in sorted(Path("results/NETWORKS").glob("*/_workflow/*.json")):
        d = json.loads(p.read_text())
        if "entries" not in d:
            continue
        entries = d["entries"]
        keys = d.get("selected_epoch_keys", list(entries))
        selected = [entries[k] for k in keys if k in entries]
        counts = dict(Counter(e["status"] for e in selected))
        if "status_counts" in d:
            record(counts == d["status_counts"], "workflow_counts", str(p))
        for e in selected:
            for f in e.get("files", []):
                artifacts += 1
                record(Path(f).is_file(), "missing_artifact", f)
            if d.get("stage") == "rayleigh":
                ray_paths.add(Path(e["update"]["rayleigh_result_file"]))
        workflows.append(
            dict(
                path=str(p),
                selected=len(keys),
                entries=len(entries),
                counts=counts,
                execution=d.get("execution_status"),
            )
        )
    print("Workflow artifacts", artifacts, flush=True)
    # Read every master-linked complete PPSD without changing any products.
    ppsd_rows = master[master.ppsd_status.eq("complete")]
    ppsd_segments = 0
    histogram_ranges = []
    source_stats = {}
    manifests = []
    for _, r in ppsd_rows.iterrows():
        p = Path(r.ppsd_npz_file)
        d = json.loads(Path(r.ppsd_provenance_file).read_text())
        s = d["summary"]
        c = d["config"]
        with np.load(p, allow_pickle=False) as a:
            t = a["_times_processed"]
            v = a["_binned_psds"]
            n = len(t)
            ppsd_segments += n
            record(str(a["id"]) == r.Seed_id, "ppsd_identity", str(p))
            record(
                n == int(r.ppsd_segments) == s["ppsd_segments"], "ppsd_count", str(p)
            )
            record(len(v) == n and np.isfinite(v).all(), "ppsd_finite", str(p))
            record(np.all(np.diff(t) > 0), "ppsd_duplicate_or_order", str(p))
            start = UTCDateTime(s["ppsd_requested_start_time"]).ns
            end = UTCDateTime(s["ppsd_requested_end_time"]).ns
            record(
                t[0] >= start - 1000
                and t[-1] + int(float(a["ppsd_length"]) * 1e9) <= end + 1000,
                "ppsd_interval",
                str(p),
            )
            record(
                float(a["sampling_rate"]) == float(r.sample_rate_hz),
                "ppsd_rate",
                str(p),
            )
            below = int(np.count_nonzero(v < float(r.ppsd_histogram_min_db)))
            above = int(np.count_nonzero(v > float(r.ppsd_histogram_max_db)))
            record(
                below + above == int(r.ppsd_histogram_out_of_range_values),
                "ppsd_histogram_count",
                str(p),
            )
            if below or above:
                histogram_ranges.append(
                    dict(
                        seed=r.Seed_id,
                        below=below,
                        above=above,
                        minimum_db=float(v.min()),
                        maximum_db=float(v.max()),
                    )
                )
        record(
            sha(r.inventory_file) == c["inventory_sha256"],
            "ppsd_source_inventory_change",
            str(p),
        )
        for item in d["source_files"]:
            source_stats[item["path"]] = (item["size"], item["mtime_ns"])
        manifests.append(
            dict(
                seed=r.Seed_id,
                origin=c["origin"],
                requested_start=s["ppsd_requested_start_time"],
                requested_end=s["ppsd_requested_end_time"],
                status=s["ppsd_status"],
            )
        )
    print("PPSD", len(ppsd_rows), "channels", ppsd_segments, "segments", flush=True)
    # Check saved array identities and closed-form arithmetic.
    # No raw reads or calibration runs.
    ray_events = 0
    candidates = 0
    spectra_count = 0
    p_s_overlaps = 0
    ray_summaries = []
    for p in sorted(ray_paths):
        d = json.loads(p.read_text())
        ev = [e for e in d["events"] if e["status"] == "candidate"]
        ray_events += len(d["events"])
        candidates += len(ev)
        record(len(ev) == d["candidate_events"], "rayleigh_count", str(p))
        if d.get("median_pressure_multiplier") is not None:
            record(
                np.isclose(
                    d["median_pressure_multiplier"],
                    np.median([e["pressure_multiplier"] for e in ev]),
                    rtol=1e-12,
                ),
                "station_median",
                str(p),
            )
        for hmap in ["source_inventory_sha256", "response_snapshot_sha256", "inputs"]:
            for name, h in d.get(hmap, {}).items():
                record(
                    Path(name).exists() and sha(name) == h, "rayleigh_input_hash", name
                )
        for item in d.get("source_files", []):
            source_stats[item["path"]] = (item["size"], item["mtime_ns"])
        for e in d["events"]:
            spec = e.get("files", {}).get("spectra")
            if not spec:
                continue
            spectra_count += 1
            with np.load(spec, allow_pickle=False) as a:
                pp = a["pressure_psd"]
                aa = a["acceleration_psd"]
                cross = a["cross_pa"]
                ratio = a["psd_ratio_pa"]
                theory = a["expected_pa"]
                expected = [
                    ("psd_ratio_pa", np.sqrt(pp / aa)),
                    ("h1_pa", cross / aa),
                    ("coherence_squared", np.clip(abs(cross) ** 2 / (pp * aa), 0, 1)),
                    ("phase_deg", np.angle(cross, deg=True)),
                ]
                for k, v in expected:
                    record(
                        np.allclose(a[k], v, rtol=1e-11, atol=1e-11),
                        "spectral_identity",
                        f"{spec}:{k}",
                    )
                valid = (
                    a["coherence_squared"]
                    >= d["settings"]["coherence_squared_threshold"]
                ) & (abs(a["phase_deg"]) <= d["settings"]["phase_limit_deg"])
                record(np.array_equal(valid, a["accepted_bins"]), "spectral_mask", spec)
                f = a["frequency_hz"]
                record(
                    np.all(np.diff(f) > 0) and np.isfinite(ratio).all(),
                    "spectral_grid",
                    spec,
                )
                if e["status"] == "candidate":
                    obs = ratio[valid]
                    ref = theory[valid]
                    g = float(np.dot(obs, ref) / np.dot(obs, obs))
                    record(
                        np.isclose(g, e["pressure_multiplier"], rtol=1e-12),
                        "candidate_gain",
                        spec,
                    )
                    record(
                        np.isclose(
                            g * e["sensitivity_counts_per_pa_multiplier"], 1, rtol=1e-12
                        ),
                        "gain_reciprocal",
                        spec,
                    )
                    record(
                        valid.sum() >= 3
                        and np.median(a["coherence_squared"])
                        >= d["settings"]["coherence_squared_threshold"]
                        and np.mean(a["coherence_squared"])
                        >= d["settings"]["coherence_squared_threshold"]
                        and not e.get("other_predicted_arrival_ids"),
                        "candidate_qc",
                        spec,
                    )
                    left, right = UTCDateTime(e["start_utc"]), UTCDateTime(e["end_utc"])
                    if any(
                        e.get(k) and left <= UTCDateTime(e[k]) <= right
                        for k in ["predicted_first_p_utc", "predicted_first_s_utc"]
                    ):
                        p_s_overlaps += 1
        ray_summaries.append(
            dict(
                path=str(p),
                station=d["station"],
                status=d["status"],
                candidate_events=len(ev),
            )
        )
    print(
        "Rayleigh",
        len(ray_paths),
        "stations/runs",
        spectra_count,
        "spectra",
        candidates,
        "candidate events",
        flush=True,
    )
    for p, (size, mtime) in source_stats.items():
        q = Path(p)
        record(
            q.exists() and (q.stat().st_size, q.stat().st_mtime_ns) == (size, mtime),
            "raw_source_changed",
            p,
        )
    # Audit the other completed diagnostics from saved arrays, never raw counts.
    diagnostics = Counter()
    pressure_rows = master[
        master.ppsd_units.eq("Pa^2/Hz") & master.ppsd_status.eq("complete")
    ]
    for _, row in pressure_rows.iterrows():
        selected = [
            name
            for name in ("pressure_expected_result_file", "infragravity_result_file")
            if row.get(name)
        ]
        if selected:
            with np.load(row.ppsd_npz_file, allow_pickle=False) as a:
                values = np.asarray(a["_binned_psds"], dtype=float)
                frequency = 1 / a["_period_binning"][2]
                percentiles = np.percentile(values, [5, 50, 95], axis=0)
            for name in selected:
                result = json.loads(Path(row[name]).read_text())
                table = pd.read_csv(result["files"]["frequency_csv"])
                for _, bin_row in table.iterrows():
                    index = int(np.argmin(abs(frequency - bin_row.frequency_hz)))
                    observed = bin_row[
                        ["p05_db_pa2_hz", "median_db_pa2_hz", "p95_db_pa2_hz"]
                    ].to_numpy(dtype=float)
                    record(
                        np.allclose(observed, percentiles[:, index], atol=1e-10),
                        "diagnostic_percentiles",
                        f"{row.Seed_id}:{name}:{bin_row.frequency_hz}",
                    )
                    if name == "infragravity_result_file":
                        median = bin_row.median_minus_model_db
                        if np.isfinite(median):
                            record(
                                np.isclose(
                                    median,
                                    bin_row.median_db_pa2_hz - bin_row.model_db_pa2_hz,
                                ),
                                "infragravity_residual",
                                row.Seed_id,
                            )
                    elif bin_row.comparison != "excluded":
                        fraction = np.mean(
                            (values[:, index] >= bin_row.reference_low_db_pa2_hz)
                            & (values[:, index] <= bin_row.reference_high_db_pa2_hz)
                        )
                        record(
                            np.isclose(fraction, bin_row.window_fraction_within_bounds),
                            "pressure_bound_fraction",
                            row.Seed_id,
                        )
                diagnostics[name] += 1
        if row.get("infragravity_pilot_result_file"):
            result = json.loads(Path(row.infragravity_pilot_result_file).read_text())
            path = result.get("files", {}).get("spectra.npz")
            if path:
                with np.load(path, allow_pickle=False) as a:
                    corrected = a["counts_psd"] / a["response_power_counts2_pa2"]
                    record(
                        np.allclose(corrected, a["pressure_psd"], rtol=1e-12),
                        "pilot_response_power",
                        row.Seed_id,
                    )
                    record(
                        np.allclose(
                            corrected[a["retained"]].mean(axis=0),
                            a["mean_retained_pa2_hz"],
                            rtol=1e-12,
                        ),
                        "pilot_linear_average",
                        row.Seed_id,
                    )
                diagnostics["raw_godin_pilot"] += 1
        if row.get("paired_result_file"):
            result = json.loads(Path(row.paired_result_file).read_text())
            path = result.get("files", {}).get("spectra")
            if path:
                with np.load(path, allow_pickle=False) as a:
                    pp = a["pressure_window_psd"].mean(axis=0)
                    zz = a["displacement_window_psd"].mean(axis=0)
                    cross = a["cross_window_dp"].mean(axis=0)
                    valid = a["valid"]
                    for name, expected in (
                        ("pressure_psd", pp),
                        ("displacement_psd", zz),
                        ("cross_dp", cross),
                        ("transfer_dp", cross / pp),
                        (
                            "coherence_squared",
                            np.clip(abs(cross) ** 2 / (pp * zz), 0, 1),
                        ),
                    ):
                        record(
                            np.allclose(a[name][valid], expected[valid], rtol=1e-12),
                            "paired_spectral_identity",
                            f"{row.Seed_id}:{name}",
                        )
                diagnostics["paired_spectra"] += 1
    print("Other diagnostic arrays", dict(diagnostics), flush=True)
    # Compare CSV/Excel values, numeric formatting and split long text.
    book = load_workbook(
        "exports/master_inventory.xlsx", read_only=True, data_only=False
    )
    sheet = book["Master inventory"]
    rows = sheet.iter_rows(values_only=True)
    headers = next(rows)
    checked = 0
    record(sheet.max_row - 1 == len(master), "master_row_count", str(sheet.max_row))
    for idx, row in enumerate(rows):
        if idx >= len(master):
            break
        vals = dict(zip(headers, row))
        saved = master.iloc[idx]
        for name in master.columns:
            a = saved[name]
            b = vals.get(name)
            b = "" if b is None else str(b)
            if len(a) > 15000:
                b += "".join(
                    str(vals.get(f"{name}_part_{j}") or "")
                    for j in range(2, 2 + (len(a) - 1) // 15000)
                )
            same = a == b
            if not same and a and b:
                try:
                    same = bool(np.isclose(float(a), float(b), rtol=1e-14, atol=1e-14))
                except ValueError:
                    pass
            record(
                same,
                "master_csv_excel_value",
                f"row {idx + 2} field {name}: CSV {a[:60]!r} Excel {b[:60]!r}",
            )
            checked += 1
    book.close()
    print("CSV/Excel cells", checked, flush=True)
    # Historical products stay immutable; master appends are expected.
    history = []
    for path in [
        Path("results/VALIDATION/YV_rayleigh/conventions_22bf4b45f61d/validation.json"),
        Path("docs/yl_all_zha_comparison.json"),
    ]:
        d = json.loads(path.read_text())
        ok = 0
        master_changed = []
        for name, h in d.get("protected_sha256", {}).items():
            q = Path(name)
            if q.exists() and sha(q) == h:
                ok += 1
            elif q.name.startswith("master_inventory.") and "YV_rayleigh" in str(path):
                master_changed.append(name)
            else:
                issues.append(dict(category="historical_protected_change", detail=name))
        history.append(
            dict(
                snapshot=str(path),
                matched=ok,
                expected_old_master_changes=master_changed,
            )
        )
    print("Historical hashes", history, flush=True)
    # Verify raw provenance stats and the full baseline; no raw-content rehash.
    changed = [
        p
        for p, s in before.items()
        if not Path(p).exists()
        or (Path(p).stat().st_size, Path(p).stat().st_mtime_ns) != s
    ]
    record(not changed, "audit_changed_production", changed[:10])
    record(
        all(sha(p) == h for p, h in protected.items()),
        "audit_changed_master_or_inventory",
        "current baseline",
    )
    report = dict(
        scope=(
            "Read-only repository, saved-array and provenance audit; "
            "no raw spectra recomputation or calibration acceptance"
        ),
        source_commit=subprocess.check_output(
            ["git", "rev-parse", "HEAD"], text=True
        ).strip(),
        tracked_files=len(tracked),
        parsed=dict(parse),
        broken_local_links=broken,
        workflow=workflows,
        ppsd_channels=len(ppsd_rows),
        ppsd_segments=ppsd_segments,
        ppsd_histogram_out_of_range=histogram_ranges,
        ppsd_intervals=manifests,
        rayleigh_runs=len(ray_paths),
        rayleigh_events=ray_events,
        rayleigh_spectra=spectra_count,
        rayleigh_candidate_events=candidates,
        candidate_windows_containing_first_body_arrivals=p_s_overlaps,
        rayleigh_summaries=ray_summaries,
        raw_provenance_files=len(source_stats),
        other_saved_diagnostics=dict(diagnostics),
        csv_excel_cells_checked=checked,
        scores_nonblank={
            c: int((master[c] != "").sum()) for c in master if c.endswith("_score")
        },
        historical_protection=history,
        baseline_stat_files=len(before),
        baseline_changed=changed,
        current_protected_sha256={str(p): h for p, h in protected.items()},
        issues=issues,
    )
    (OUT / "audit.json").write_text(json.dumps(report, indent=2) + "\n")
    print(
        "ISSUES", len(issues), dict(Counter(x["category"] for x in issues)), flush=True
    )
    print("DETAILS", issues[:12], flush=True)
    print("REPORT", OUT, flush=True)
    raise SystemExit(bool(issues or broken))


if __name__ == "__main__":
    main()
