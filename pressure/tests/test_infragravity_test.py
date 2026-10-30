"""Verify Godin normalization, dispersion, smoothing and diagnostic preservation."""

import hashlib
from types import SimpleNamespace

import matplotlib.pyplot as plt
import numpy as np
import pytest
from openpyxl import load_workbook

from compy_pressure.infragravity_test import (
    GRAVITY,
    SCORE_COLUMN,
    STATUS_COLUMN,
    compare_spectra,
    comparison_plots,
    dispersion_kh,
    godin_pressure_db,
    run_infragravity_test,
)
from compy_pressure.inventory import (
    load_master_dataset,
    update_infragravity_test_results,
)
from compy_pressure.pressure_test import response_band


def test_godin_equation_at_known_dimensionless_wavenumbers():
    # Construct frequencies from chosen kH, so roots and pressure normalization
    # are known independently of the numerical dispersion solver.
    x = np.array([0.001, 0.1, 1, 3, 20])
    depth = 2918
    frequencies = np.sqrt(GRAVITY * x * np.tanh(x) / depth) / (2 * np.pi)
    np.testing.assert_allclose(dispersion_kh(frequencies, depth), x, rtol=1e-12)
    expected_power = 2000 * np.tanh(x) ** -1.5 / np.cosh(x) ** 2
    np.testing.assert_allclose(
        godin_pressure_db(frequencies, depth), 10 * np.log10(expected_power), atol=1e-10
    )
    np.testing.assert_allclose(
        godin_pressure_db(frequencies, depth, tanh_factor=False),
        10 * np.log10(2000 / np.cosh(x) ** 2),
        atol=1e-10,
    )
    # At a fixed frequency, deeper water must attenuate seafloor pressure.
    assert godin_pressure_db(0.01, 4000) < godin_pressure_db(0.01, 2000)
    assert np.isfinite(godin_pressure_db(1, 10000))


@pytest.mark.parametrize(
    "frequency,depth", [(0, 1000), (-0.01, 1000), (np.nan, 1000), (0.01, 0), (0.01, "")]
)
def test_invalid_model_inputs(frequency, depth):
    with pytest.raises(ValueError, match="finite and positive"):
        godin_pressure_db(frequency, depth)


def test_bin_averaging_band_selection_and_no_extrapolation():
    ppsd = SimpleNamespace(
        period_bin_centers=1 / np.array([0.005, 0.015, 0.025]),
        period_bin_right_edges=1 / np.array([0.004, 0.01, 0.02]),
        period_bin_left_edges=1 / np.array([0.008, 0.02, 0.03]),
        sampling_rate=1,
        nfft=1000,
    )
    band = {"minimum_hz": 0.009, "maximum_hz": 0.1}
    values = np.array([[30, 40, 50], [32, 42, 52], [34, 44, 54]], dtype=float)
    rows, _ = compare_spectra(ppsd, values, 2918, band)
    assert [r["comparison"] for r in rows] == ["excluded", "compared", "excluded"]
    assert rows[0]["exclusion_reason"] == "smoothing_bin_outside_sensor_band"
    assert rows[0]["median_minus_model_db"] is None
    assert rows[2]["model_db_pa2_hz"] is None
    expected = np.mean(godin_pressure_db(np.arange(10, 21) / 1000, 2918))
    assert rows[1]["model_db_pa2_hz"] == pytest.approx(expected)
    assert rows[1]["median_minus_model_db"] == pytest.approx(42 - expected)
    # Nonlinear steep spectra must be averaged over bin frequencies, not just
    # evaluated at the center; dB averaging matches the stored ObsPy smoother.
    assert abs(expected - godin_pressure_db(0.015, 2918)) > 0.1
    with pytest.raises(ValueError, match="Nyquist"):
        compare_spectra(ppsd, values, 2918, band, model_max=1)


def test_infragravity_run_preserves_other_tests_and_master_data(saved_pressure):
    master_dir, source, _ = saved_pressure
    before = load_master_dataset(master_dir)
    source_hash = hashlib.sha256(source.read_bytes()).hexdigest()
    result = run_infragravity_test(
        master_dir, network="XX", station="TEST", show=False
    )[0]
    after = load_master_dataset(master_dir)
    assert result["status"] == "comparison_only"
    assert result["score"] is None
    assert result["depth_m"] == 1000
    assert result["depth_basis"] == "inventory_elevation_estimate"
    assert result["compared_bins"] > 0
    assert after[before.columns.drop(STATUS_COLUMN)].equals(
        before.drop(columns=STATUS_COLUMN)
    )
    assert after.loc[0, SCORE_COLUMN] == ""
    assert after.loc[0, STATUS_COLUMN] == "comparison_only"
    assert after.loc[1, "infragravity_status"] == ""
    assert hashlib.sha256(source.read_bytes()).hexdigest() == source_hash
    book = load_workbook(master_dir / "master_inventory.xlsx")
    assert book["Manual notes"]["A1"].value == "preserve me"
    sheet = book["Master inventory"]
    headers = {c.value: c.column for c in sheet[1]}
    assert sheet.cell(
        2, headers["infragravity_median_residual_db"]
    ).value == pytest.approx(result["median_residual_db"])
    assert sheet.cell(2, headers[SCORE_COLUMN]).value is None
    book.close()
    row = after.iloc[0]
    with pytest.raises(ValueError, match="Updates may only"):
        update_infragravity_test_results(
            master_dir,
            [
                {
                    **row[
                        [
                            "Seed_id",
                            "channel_start_time",
                            "channel_end_time",
                            "inventory_file",
                        ]
                    ],
                    SCORE_COLUMN: 1,
                }
            ],
        )


def test_empty_overlap_is_recorded_and_depth_override_is_explicit(saved_pressure):
    master_dir, _, _ = saved_pressure
    result = run_infragravity_test(
        master_dir,
        network="XX",
        station="TEST",
        show=False,
        depth_m=2918,
        sensor_min=0.5,
        sensor_max=3,
    )[0]
    assert result["status"] == "no_comparable_bins"
    assert result["compared_bins"] == 0
    assert result["median_residual_db"] is None
    assert result["effective_band_min_hz"] is None
    assert result["effective_band_max_hz"] is None
    assert result["depth_basis"] == "user_supplied"
    assert not result["depth_independently_verified"]


def test_godin_plot_coordinates_match_diagnostics(saved_pressure):
    _, source, ppsd = saved_pressure
    band, _, _ = response_band(ppsd, source.parent / "inventory.xml")
    rows, percentiles = compare_spectra(ppsd, np.asarray(ppsd.psd_values), 1000, band)
    figures = comparison_plots(ppsd, percentiles, rows, 1000, band)
    assert len(figures) == 3
    model_lines = [
        line
        for line in figures[1].axes[0].lines
        if line.get_label() == "Godin averaged over PSD bins"
    ]
    available = [row for row in rows if row["model_db_pa2_hz"] is not None]
    np.testing.assert_allclose(
        model_lines[0].get_xdata(), [r["frequency_hz"] for r in available]
    )
    np.testing.assert_allclose(
        model_lines[0].get_ydata(), [r["model_db_pa2_hz"] for r in available]
    )
    included = [row for row in rows if row["comparison"] == "compared"]
    detail = figures[2].axes[1]
    medians = next(
        line for line in detail.lines if line.get_label() == "Eligible median residual"
    )
    np.testing.assert_allclose(
        medians.get_ydata(), [r["median_minus_model_db"] for r in included]
    )
    np.testing.assert_allclose(
        medians.get_xdata(), [r["frequency_hz"] * 1000 for r in included]
    )
    support = detail.collections[0].get_segments()
    spread = detail.collections[1].get_segments()
    for index, row in enumerate(included):
        np.testing.assert_allclose(
            support[index][:, 0],
            [row["smoothing_minimum_hz"] * 1000, row["smoothing_maximum_hz"] * 1000],
        )
        np.testing.assert_allclose(
            spread[index][:, 1],
            [
                row["p05_db_pa2_hz"] - row["model_db_pa2_hz"],
                row["p95_db_pa2_hz"] - row["model_db_pa2_hz"],
            ],
        )
    diagnostics = next(
        line
        for line in figures[2].axes[0].lines
        if line.get_label() == "Diagnostic median"
    )
    np.testing.assert_allclose(
        diagnostics.get_ydata(),
        [row["median_db_pa2_hz"] - row["model_db_pa2_hz"] for row in available],
    )
    for figure in figures:
        plt.close(figure)


def test_independent_dispersion_solver_and_exact_obspy_mask(saved_pressure):
    from scipy.optimize import brentq

    _, source, ppsd = saved_pressure
    band, _, _ = response_band(ppsd, source.parent / "inventory.xml")
    rows, _ = compare_spectra(ppsd, np.asarray(ppsd.psd_values), 1000, band)
    for row in rows:
        if row["model_db_pa2_hz"] is None:
            continue
        index = int(np.argmin(abs(1 / ppsd.period_bin_centers - row["frequency_hz"])))
        mask = (ppsd.psd_periods >= ppsd.period_bin_left_edges[index]) & (
            ppsd.psd_periods <= ppsd.period_bin_right_edges[index]
        )
        independent_values = []
        for frequency in 1 / ppsd.psd_periods[mask]:
            a = (2 * np.pi * frequency) ** 2 * 1000 / GRAVITY
            x = brentq(
                lambda value: value * np.tanh(value) - a,
                0,
                a + np.sqrt(a) + 1,
                xtol=1e-14,
            )
            independent_values.append(
                10 * np.log10(2000 * np.tanh(x) ** (-1.5) / np.cosh(x) ** 2)
            )
        assert row["model_db_pa2_hz"] == pytest.approx(
            np.mean(independent_values), abs=1e-10
        )


def test_narrow_model_band_without_saved_centers_is_not_a_failure(saved_pressure):
    master_dir, _, _ = saved_pressure
    result = run_infragravity_test(
        master_dir,
        network="XX",
        station="TEST",
        show=False,
        model_min=0.012,
        model_max=0.01201,
    )[0]
    assert result["status"] == "no_comparable_bins"
    assert result["diagnostic_bins"] == 0
    assert result["median_residual_db"] is None
    assert result["score"] is None
