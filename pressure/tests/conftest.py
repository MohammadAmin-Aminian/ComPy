# Moved from BRUIT-FM 1.3.0.
import sys

sys.modules["bruitfm"] = None  # the moved code must not import BRUIT-FM

"""Saved pressure data shared by pressure and infragravity integration tests."""

import numpy as np
import pytest
from obspy import Trace, UTCDateTime
from obspy.core.inventory import Channel, Inventory, Network, Site, Station
from obspy.core.inventory.response import Response
from obspy.signal import PPSD
from openpyxl import load_workbook

from compy_pressure.inventory import (
    TEST_COLUMNS,
    build_master_dataset,
    save_master_dataset,
)


@pytest.fixture
def saved_pressure(tmp_path):
    start = UTCDateTime("2020-01-01")
    response = Response.from_paz(
        zeros=[], poles=[], stage_gain=4, input_units="PA", output_units="COUNTS"
    )
    response.instrument_sensitivity.frequency = 1
    channels = [
        Channel(
            code=code,
            location_code="00",
            latitude=1,
            longitude=2,
            elevation=-1000,
            depth=0,
            sample_rate=8,
            start_date=start,
            end_date=start + 4096,
            response=response,
        )
        for code in ("BDH", "BHZ")
    ]
    inv = Inventory(
        [
            Network(
                code="XX",
                stations=[
                    Station(
                        code="TEST",
                        latitude=1,
                        longitude=2,
                        elevation=-1000,
                        site=Site(name="Synthetic"),
                        channels=channels,
                    )
                ],
            )
        ],
        source="Test",
    )
    inv_dir = tmp_path / "inventories"
    inv_dir.mkdir()
    inv.write(str(inv_dir / "source.xml"), format="STATIONXML")
    master = build_master_dataset(inv_dir).astype(object).fillna("")
    result_dir = tmp_path / "results"
    result_dir.mkdir()
    inv.select(channel="BDH").write(
        str(result_dir / "inventory.xml"), format="STATIONXML"
    )
    trace = Trace(np.random.default_rng(7).integers(-1000, 1000, 16384, dtype=np.int32))
    trace.stats.update(
        dict(
            network="XX",
            station="TEST",
            location="00",
            channel="BDH",
            sampling_rate=8,
            starttime=start,
        )
    )
    ppsd = PPSD(
        trace.stats,
        metadata=inv,
        special_handling="hydrophone",
        ppsd_length=512,
        period_limits=(0.25, 100),
        db_bins=(-200, 100, 1),
    )
    ppsd.add(trace)
    source = result_dir / "ppsd.npz"
    ppsd.save_npz(str(source))
    for index in master.index:
        master.loc[index, "ppsd_units"] = "Pa^2/Hz" if index == 0 else "m^2/s^4/Hz"
        master.loc[index, "ppsd_status"] = "complete"
        master.loc[index, "ppsd_npz_file"] = str(source)
    master.loc[1, TEST_COLUMNS[0]] = 0.75
    export_dir = tmp_path / "exports"
    save_master_dataset(master, export_dir)
    book = load_workbook(export_dir / "master_inventory.xlsx")
    book.create_sheet("Manual notes")["A1"] = "preserve me"
    book.save(export_dir / "master_inventory.xlsx")
    book.close()
    return export_dir, source, ppsd
