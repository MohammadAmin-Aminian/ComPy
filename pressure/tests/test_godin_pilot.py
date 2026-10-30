"""Verify full-window PSD normalization, screening and safe pilot integration."""

import hashlib
from pathlib import Path

import numpy as np
import pytest
from obspy import Stream, Trace, UTCDateTime
from openpyxl import load_workbook

from compy_pressure import godin_pilot as pilot
from compy_pressure.inventory import TEST_COLUMNS, load_master_dataset


def test_daily_schedule_has_three_full_hann_windows():
    start = UTCDateTime("2020-01-01")
    windows = pilot.daily_windows(start, 7)
    assert len(windows) == 21
    np.testing.assert_allclose([s - start for s, _ in windows[:3]], [0, 20844, 41688])
    assert all(end - begin == 44712 for begin, end in windows)
    assert windows[2][1] == start + 86400
    assert windows[-1][1] == start + 7 * 86400
    with pytest.raises(ValueError):
        pilot.daily_windows(start, 0)


@pytest.mark.parametrize("npts", [1000, 1001])
def test_periodogram_matches_independent_full_fft_and_parseval(npts):
    fs = 10.0
    data = np.random.default_rng(2).normal(size=npts) + 500
    frequency, actual = pilot.count_spectrum(data, fs, fs / 2)
    hann = 0.5 - 0.5 * np.cos(2 * np.pi * np.arange(npts) / npts)
    tapered = (data - data.mean()) * hann
    expected = abs(np.fft.rfft(tapered)) ** 2 / (fs * np.sum(hann**2))
    expected[1:] *= 2
    if npts % 2 == 0:
        expected[-1] /= 2
    np.testing.assert_allclose(actual, expected[1:], rtol=1e-12, atol=1e-14)
    np.testing.assert_allclose(frequency, np.fft.rfftfreq(npts, 1 / fs)[1:])
    assert (expected.sum() * fs / npts) == pytest.approx(
        np.sum(tapered**2) / np.sum(hann**2), rel=1e-12
    )


def test_screening_rejects_one_of_21_and_keeps_threshold_ties():
    kept, threshold = pilot.screen_windows(np.arange(21))
    assert threshold == 19
    assert kept.sum() == 20
    assert not kept[-1]
    kept, _ = pilot.screen_windows(np.ones(21))
    assert kept.all()
    with pytest.raises(ValueError):
        pilot.screen_windows([np.nan])


def test_continuous_window_alignment_gaps_and_flat_data():
    start = UTCDateTime("2020-01-01")
    trace = Trace(np.arange(100, dtype=float))
    trace.stats.starttime = start + 0.004
    trace.stats.sampling_rate = 10
    data, actual = pilot.complete_samples(Stream([trace]), start, 10, 50)
    assert actual == start + 0.004
    assert len(data) == 50
    assert pilot.complete_samples(Stream([trace]), start, 10, 101) == (None, None)
    trace.data = np.ones(100)
    assert pilot.complete_samples(Stream([trace]), start, 10, 50) == (None, None)
    trace.data = np.ma.array(np.arange(100), mask=np.arange(100) == 20)
    assert pilot.complete_samples(Stream([trace]), start, 10, 50) == (None, None)


def test_real_mseed_pilot_response_linear_mean_and_master_preservation(
    saved_pressure, monkeypatch
):
    master_dir, source, _ = saved_pressure
    monkeypatch.setattr(pilot, "WINDOW_SECONDS", 512)
    monkeypatch.setattr(pilot, "DAY_SECONDS", 2048)
    data_dir = master_dir.parent / "raw"
    data_dir.mkdir()
    trace = Trace(
        np.random.default_rng(55).integers(-1000, 1000, 16384, dtype=np.int32)
    )
    trace.stats.update(
        dict(
            network="XX",
            station="TEST",
            location="00",
            channel="BDH",
            sampling_rate=8,
            starttime=UTCDateTime("2020-01-01"),
        )
    )
    raw = data_dir / "XX.TEST.00.BDH.D.2020.001"
    trace.write(str(raw), format="MSEED")
    before = load_master_dataset(master_dir)
    hashes = {p: hashlib.sha256(p.read_bytes()).hexdigest() for p in [source, raw]}
    summary = pilot.run_godin_pilot(
        master_dir,
        data_dir=data_dir,
        network="XX",
        station="TEST",
        days=1,
        model_min=0.005,
        model_max=0.02,
        show=False,
    )
    assert summary["requested_windows"] == summary["complete_windows"] == 3
    assert summary["retained_windows"] == 2
    assert summary["score"] is None
    npz = np.load(summary["files"]["spectra.npz"])
    np.testing.assert_allclose(npz["pressure_psd"], npz["counts_psd"] / 16, rtol=1e-10)
    np.testing.assert_allclose(
        npz["mean_retained_pa2_hz"], npz["pressure_psd"][npz["retained"]].mean(axis=0)
    )
    after = load_master_dataset(master_dir)
    assert after[before.columns].equals(before)
    assert after.loc[0, "infragravity_pilot_retained_windows"] == "2"
    assert after[TEST_COLUMNS].equals(before[TEST_COLUMNS])
    for p, digest in hashes.items():
        assert hashlib.sha256(p.read_bytes()).hexdigest() == digest
    book = load_workbook(master_dir / "master_inventory.xlsx")
    assert book["Manual notes"]["A1"].value == "preserve me"
    headers = [c.value for c in book["Master inventory"][1]]
    column = headers.index("infragravity_pilot_retained_windows") + 1
    assert book["Master inventory"].cell(2, column).value == 2
    book.close()
    for name in ["pressure_model.png", "screening.png", "ppsd_comparison.png"]:
        assert Path(summary["files"][name]).stat().st_size > 1000
    master_hash = hashlib.sha256(
        (master_dir / "master_inventory.csv").read_bytes()
    ).hexdigest()
    assert pilot.show_saved_pilot(master_dir, "XX", "TEST", show=False) == summary
    assert (
        hashlib.sha256((master_dir / "master_inventory.csv").read_bytes()).hexdigest()
        == master_hash
    )


def test_middle_interval_does_not_use_deployment_start(saved_pressure, monkeypatch):
    master_dir, source, _ = saved_pressure
    # Stop after interval selection: 329-day coverage should choose an interior week.
    from compy_pressure.archive import WaveformFile

    start = UTCDateTime("2020-01-01")
    record = WaveformFile(source, "XX.TEST.00.BDH", start, start + 329 * 86400, 8)
    monkeypatch.setattr(pilot, "discover_waveforms", lambda *args: [record])
    # Extend the synthetic channel epoch to cover this scheduling-only check.
    real_select = pilot.select_epoch

    def select(row):
        inv, channel = real_select(row)
        channel.end_date = start + 330 * 86400
        return inv, channel

    monkeypatch.setattr(pilot, "select_epoch", select)
    seen = []

    def stop(records, begin, end):
        seen.append(begin)
        raise RuntimeError("stop after scheduling")

    monkeypatch.setattr(pilot, "read_window", stop)
    with pytest.raises(RuntimeError, match="stop after scheduling"):
        pilot.run_godin_pilot(
            master_dir, network="XX", station="TEST", days=7, show=False
        )
    assert seen[0] >= start + 160 * 86400
    assert seen[0] < start + 163 * 86400
