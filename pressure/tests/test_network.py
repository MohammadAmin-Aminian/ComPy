"""Exercise discovery, multi-epoch batches, restart and inventory preservation."""

import argparse
import hashlib
import os
from pathlib import Path
from threading import Event

import numpy as np
import pytest
from obspy import Trace, UTCDateTime
from obspy.core.inventory import Channel, Inventory, Network, Site, Station
from obspy.core.inventory.response import Response
from openpyxl import load_workbook

from compy_pressure import network
from compy_pressure.archive import index_waveform_files
from compy_pressure.inventory import (
    RESULT_COLUMNS,
    load_master_dataset,
    synchronize_inventory,
)
from compy_pressure.ppsd import process_epoch

try:  # campaign-only stage; absent from the core exported to ComPy
    from compy_pressure import godin_pilot
except ImportError:
    godin_pilot = None


@pytest.fixture
def network_data(tmp_path):
    start = UTCDateTime("2020-01-01")
    inventories, data, exports = [
        tmp_path / name for name in ("inventories", "Data", "exports")
    ]
    inventories.mkdir()

    def channel(code, location="00", offset=0, length=4096):
        response = Response.from_paz(
            zeros=[],
            poles=[],
            stage_gain=4,
            input_units="M/S" if code == "BHZ" else "PA",
            output_units="COUNTS",
        )
        response.instrument_sensitivity.frequency = 1
        return Channel(
            code=code,
            location_code=location,
            latitude=1,
            longitude=2,
            elevation=-1000,
            depth=0,
            sample_rate=8,
            start_date=start + offset,
            end_date=start + offset + length,
            response=response,
        )

    for net, sta, channels in [
        (
            "XX",
            "TEST",
            [
                channel("BDH", length=2048),
                channel("BDH", offset=2048, length=2048),
                channel("BDH", location="10"),
                channel("BHZ"),
            ],
        ),
        ("ZZ", "OTHER", [channel("BDH")]),
    ]:
        station = Station(
            code=sta,
            latitude=1,
            longitude=2,
            elevation=-1000,
            site=Site(name=sta),
            channels=channels,
        )
        Inventory([Network(code=net, stations=[station])], source="Test").write(
            str(inventories / f"{net}.xml"), format="STATIONXML"
        )
        for location, code in sorted({(c.location_code, c.code) for c in channels}):
            trace = Trace(
                np.random.default_rng(9).integers(-1000, 1000, 32768, dtype=np.int32)
            )
            trace.stats.update(
                dict(
                    network=net,
                    station=sta,
                    location=location,
                    channel=code,
                    sampling_rate=8,
                    starttime=start,
                )
            )
            folder = data / net / sta / f"{code}.D"
            folder.mkdir(parents=True, exist_ok=True)
            trace.write(
                str(folder / f"{net}.{sta}.{location}.{code}.D.2020.001"),
                format="MSEED",
            )
    synchronize_inventory(inventories, exports)
    book = load_workbook(exports / "master_inventory.xlsx")
    book.create_sheet("Manual notes")["A1"] = "preserve me"
    book.save(exports / "master_inventory.xlsx")
    book.close()
    return inventories, data, exports, tmp_path / "results"


def test_filename_discovery_all_networks_and_filters(network_data):
    _, data, exports, _ = network_data
    files = network.scan_files(data)
    assert len(files) == 4
    assert set(network.scan_files(data, networks=["ZZ"])) == {"ZZ.OTHER.00.BDH"}
    assert len(network.scan_files(data, stations=["TEST"])) == 3
    plan = network.build_plan(load_master_dataset(exports), files, data)
    assert plan["networks"] == ["XX", "ZZ"]
    assert plan["channel_epochs"] == 5
    assert plan["station_count"] == 2
    assert plan["unmatched_seed_ids"] == []
    assert len(index_waveform_files(files["XX.TEST.00.BDH"])) == 1
    with pytest.raises(ValueError, match="No matching"):
        network.scan_files(data, networks=["MISSING"])


def test_full_network_stages_multiple_locations_epochs_and_cache(
    network_data, monkeypatch
):
    _, data, exports, results = network_data
    files = network.scan_files(data)
    original = load_master_dataset(exports)
    hashes = {
        p: hashlib.sha256(p.read_bytes()).hexdigest()
        for paths in files.values()
        for p in paths
    }
    options = dict(ppsd_length=512, minimum_frequency=0.02, chunk_seconds=1024)
    report = network.run_stage(
        "ppsd", exports, results, files, workers=2, options=options
    )
    assert report["status_counts"] == {"complete": 5}
    saved = load_master_dataset(exports)
    assert saved.ppsd_status.eq("complete").all()
    assert set(saved.ppsd_units) == {"Pa^2/Hz", "m^2/s^4/Hz"}

    def unexpected_job(*args):
        raise AssertionError("Completed PPSD should have been reused")

    monkeypatch.setattr(network, "ppsd_job", unexpected_job)
    reused = network.run_stage(
        "ppsd", exports, results, files, workers=1, options=options
    )
    assert reused["status_counts"] == {"complete": 5}
    stages = ["plots", "pressure", "infragravity"]
    if godin_pilot is not None:
        monkeypatch.setattr(godin_pilot, "WINDOW_SECONDS", 512)
        monkeypatch.setattr(godin_pilot, "DAY_SECONDS", 2048)
        stages.append("pilot")
    for stage in stages:
        report = network.run_stage(
            stage,
            exports,
            results,
            files,
            options={"days": 1, "start": None} if stage == "pilot" else {},
        )
        assert sum(report["status_counts"].values()) == (5 if stage == "plots" else 4)
        assert set(report["status_counts"]).issubset(network.SUCCESS)
        for key in report["selected_epoch_keys"]:
            assert all(Path(p).is_file() for p in report["entries"][key]["files"])
    after = load_master_dataset(exports)
    assert after[RESULT_COLUMNS].equals(original[RESULT_COLUMNS])
    assert after.Seed_id.tolist() == original.Seed_id.tolist()
    if godin_pilot is not None:
        assert (
            after.loc[after.input_units.eq("PA"), "infragravity_pilot_status"]
            .eq("comparison_only")
            .all()
        )
        assert (
            after.loc[after.input_units.eq("M/S"), "infragravity_pilot_status"]
            .eq("")
            .all()
        )
    for path, digest in hashes.items():
        assert hashlib.sha256(path.read_bytes()).hexdigest() == digest
    book = load_workbook(exports / "master_inventory.xlsx")
    assert book["Manual notes"]["A1"].value == "preserve me"
    book.close()
    selected = network.scan_files(data, networks=["ZZ"])
    if godin_pilot is not None:
        assert len(network.view_plots(exports, selected, "pilot", show=False)) == 3


def test_pressure_only_ppsd_runs_largest_channels_first(network_data, monkeypatch):
    _, data, exports, results = network_data
    files = network.scan_files(data)
    big = files["XX.TEST.10.BDH"][0]
    big.write_bytes(big.read_bytes() + b"\0" * 100000)
    order = []

    def fake_job(row, paths, results_directory, options):
        order.append(row["Seed_id"])
        return {n: row[n] for n in network.KEYS} | {"ppsd_status": "failed"}

    monkeypatch.setattr(network, "ppsd_job", fake_job)
    report = network.run_stage(
        "ppsd", exports, results, files, workers=1, pressure_only=True
    )
    assert report["status_counts"] == {"failed": 4}
    assert "XX.TEST.00.BHZ" not in order
    assert order[0] == "XX.TEST.10.BDH"
    assert network.worker_count("max") == max(1, (os.cpu_count() or 2) - 1)
    assert network.worker_count("3") == 3
    with pytest.raises(argparse.ArgumentTypeError):
        network.worker_count("0")


def test_channel_patterns_select_channel_codes(network_data, monkeypatch):
    _, data, exports, results = network_data
    files = network.scan_files(data)
    order = []

    def fake_job(row, paths, results_directory, options):
        order.append(row["Seed_id"])
        return {n: row[n] for n in network.KEYS} | {"ppsd_status": "failed"}

    monkeypatch.setattr(network, "ppsd_job", fake_job)
    network.run_stage("ppsd", exports, results, files, workers=1, channels=["?HZ"])
    assert order and all(seed.endswith("HZ") for seed in order)


def test_bad_channel_does_not_prevent_other_channels_and_retry_is_clean(network_data):
    _, data, exports, results = network_data
    files = network.scan_files(data, networks=["XX"])
    broken = files["XX.TEST.00.BHZ"][0]
    original = broken.read_bytes()
    broken.write_bytes(b"broken miniSEED")
    options = dict(ppsd_length=512, minimum_frequency=0.02, chunk_seconds=1024)
    report = network.run_stage(
        "ppsd", exports, results, files, workers=1, options=options
    )
    assert report["status_counts"] == {"complete": 3, "failed": 1}
    assert (
        load_master_dataset(exports)
        .loc[lambda m: m.Seed_id.eq("XX.TEST.00.BHZ"), "ppsd_status"]
        .eq("failed")
        .all()
    )
    broken.write_bytes(original)
    retry = network.run_stage(
        "ppsd", exports, results, files, workers=1, options=options
    )
    assert retry["status_counts"] == {"complete": 4}


def test_inventory_append_preserves_processed_fields_and_extra_sheet(network_data):
    inventories, _, exports, _ = network_data
    before = load_master_dataset(exports)
    csv_hash = hashlib.sha256(
        (exports / "master_inventory.csv").read_bytes()
    ).hexdigest()
    assert synchronize_inventory(inventories, exports)["added_rows"] == 0
    assert (
        hashlib.sha256((exports / "master_inventory.csv").read_bytes()).hexdigest()
        == csv_hash
    )
    from obspy import read_inventory

    inv = read_inventory(str(inventories / "ZZ.xml"))
    inv[0].code = "NEW"
    inv.write(str(inventories / "new.xml"), format="STATIONXML")
    result = synchronize_inventory(inventories, exports)
    assert result["added_rows"] == 1
    after = load_master_dataset(exports)
    assert after.iloc[: len(before)][before.columns].equals(before)
    assert after.iloc[-1].Seed_id == "NEW.OTHER.00.BDH"
    book = load_workbook(exports / "master_inventory.xlsx")
    assert book["Manual notes"]["A1"].value == "preserve me"
    book.close()


def test_cooperative_stop_saves_checkpoint_and_resumes(network_data):
    _, data, exports, results = network_data
    row = load_master_dataset(exports).iloc[0]
    files = network.scan_files(data, seed_ids=[row.Seed_id])
    records = index_waveform_files(files[row.Seed_id])
    event = Event()
    event.set()
    options = dict(ppsd_length=512, minimum_frequency=0.02, chunk_seconds=1024)
    stopped = process_epoch(row, records, results, cancel_event=event, **options)
    assert stopped["ppsd_status"] == "partial"
    assert stopped["ppsd_segments"] == 0
    resumed = process_epoch(row, records, results, **options)
    assert resumed["ppsd_status"] == "complete"
    assert resumed["ppsd_segments"] == 7


def test_stage_without_ppsd_records_failure_and_unmatched_inventory_is_explicit(
    network_data,
):
    _, data, exports, results = network_data
    files = network.scan_files(data, networks=["ZZ"])
    report = network.run_stage("pressure", exports, results, files)
    assert report["status_counts"] == {"failed": 1}
    entry = next(iter(report["entries"].values()))
    assert "run the ppsd stage" in entry["error"]
    bad_files = {"NO.UNKNOWN.00.BDH": next(iter(files.values()))}
    plan = network.build_plan(load_master_dataset(exports), bad_files, data)
    assert plan["unmatched_seed_ids"] == ["NO.UNKNOWN.00.BDH"]
    with pytest.raises(ValueError, match="Missing StationXML"):
        network.run_stage("ppsd", exports, results, bad_files)


def test_corrupt_inventory_append_leaves_both_master_files_unchanged(network_data):
    inventories, _, exports, _ = network_data
    hashes = {
        p: hashlib.sha256(p.read_bytes()).hexdigest()
        for p in exports.glob("master_inventory.*")
    }
    (inventories / "broken.xml").write_text("not StationXML")
    with pytest.raises(RuntimeError, match="broken.xml"):
        synchronize_inventory(inventories, exports)
    for path, digest in hashes.items():
        assert hashlib.sha256(path.read_bytes()).hexdigest() == digest
