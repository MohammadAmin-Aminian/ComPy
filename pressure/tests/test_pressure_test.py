"""Check reference units, band exclusion and pressure-only master updates."""

import csv
import hashlib
import json
from importlib.resources import files
from pathlib import Path
from types import SimpleNamespace

import matplotlib.pyplot as plt
import numpy as np
import pytest
from obspy import read_inventory
from obspy.core.inventory.response import Response
from openpyxl import load_workbook

from compy_pressure.inventory import (
    TEST_COLUMNS,
    load_master_dataset,
    update_pressure_test_results,
)
from compy_pressure.pressure_test import (
    REFERENCE_COLORS,
    load_references,
    ppsd_limits_plot,
    response_band,
    run_pressure_test,
    summarize,
    validate_pressure,
)
from compy_pressure.pressure_test import (
    test_details_plot as make_test_details_plot,
)


def test_digitized_pressure_units_and_attribution():
    metadata, brown = load_references()
    # Reconstruct a point from the source raster's axis coordinates: the 120 dB
    # conversion must apply to power density, not an amplitude/PSD confusion.
    expected_low = 190 - (469.5 - 15) * 150 / 708 - 120
    assert brown[np.argmin(abs(brown[:, 0] - 1)), 1] == pytest.approx(expected_low)
    assert brown[0, 0] < 0.02 and brown[-1, 0] > 70
    assert metadata["webb"]["verified_published_bounds"]
    assert "Figure 2" in metadata["webb"]["citation"]
    assert max(p["frequency_hz"] for p in metadata["webb"]["points"]) == 0.2
    assert metadata["brown"]["minimum_comparison_hz"] == 0.2
    assert np.all(brown[:, 1] < brown[:, 2])
    assert np.all(np.diff(brown[:, 0]) > 0)
    assert metadata["brown"]["conversion_to_db_pa2_hz"] == -120
    assert "not official" in metadata["brown"]["provenance"]


def test_all_full_band_samples_follow_recorded_raster_axes():
    path = Path(str(files("compy_pressure").joinpath("reference_data")))
    with (path / "brown_hydroacoustic_fullband.csv").open() as stream:
        rows = list(csv.DictReader(stream))
    assert len(rows) == 37
    for row in rows:
        x = float(row["image_x_px"])
        assert float(row["frequency_hz"]) == pytest.approx(10 ** (-2 + (x - 123) / 177))
        for bound in ("low", "high"):
            y = float(row[f"{bound}_y_px"])
            expected = 190 - (y - 15) * 150 / 708 - 120
            assert float(row[f"{bound}_db_pa2_hz"]) == pytest.approx(expected)


def test_webb_below_and_brown_above_the_split_by_default():
    metadata, brown = load_references()
    f = np.array([0.05, 0.1, 0.2, 1.0, 2.0])
    values = np.zeros((3, len(f)))
    band = {"minimum_hz": 0.01, "maximum_hz": 5, "fft_frequency_step_hz": 0.001}
    rows, _, _ = summarize(
        values, f, f / np.sqrt(2), f * np.sqrt(2), band, metadata, brown
    )
    webb = {
        r["reference_frequency_hz"]: r
        for r in rows
        if r["reference"] == "webb_instructions"
    }
    assert sorted(webb) == [0.005, 0.05, 0.1, 0.2]
    # 0.005 Hz has no saved bin here; the others are compared.
    assert webb[0.005]["exclusion_reason"] == "reference_point_outside_saved_bins"
    assert all(webb[f]["comparison"] != "excluded" for f in (0.05, 0.1, 0.2))
    brown_rows = [r for r in rows if r["reference"] == "brown_figure8"]
    assert brown_rows and all(r["frequency_hz"] >= 0.2 for r in brown_rows)
    # The 0.2 Hz bin is judged by Webb only, never by both sources.
    compared = [r["frequency_hz"] for r in rows if r["comparison"] != "excluded"]
    assert len(compared) == len(set(compared))


def test_reference_exclusion_percentiles_and_frequency_smoothing():
    frequencies = np.array([0.05, 0.1, 0.2, 1.0, 2.0, 4.0, 8.0])
    values = np.tile([-20, 0, 40, -5, -15, -50, -20], (3, 1)).astype(float)
    values += np.array([-1, 0, 1])[:, None]
    metadata, _ = load_references()
    # Constant boundaries make the expected smoothed bounds independently known.
    brown = np.array([[1, -30, -10], [5, -30, -10]], dtype=float)
    band = {"minimum_hz": 0.03, "maximum_hz": 5, "fft_frequency_step_hz": 0.01}
    rows, percentiles, fraction = summarize(
        values,
        frequencies,
        frequencies / np.sqrt(2),
        frequencies * np.sqrt(2),
        band,
        metadata,
        brown,
    )
    included = [r for r in rows if r["comparison"] != "excluded"]
    # Webb at 0.05, 0.1 and 0.2 Hz; Brown only at 2 Hz (the 1 Hz bin starts
    # below this test curve's 1 Hz support).
    assert [r["comparison"] for r in included] == ["within"] * 4
    assert fraction == 1.0
    assert np.all(percentiles[0] <= percentiles[1])
    assert np.all(percentiles[1] <= percentiles[2])
    tested_brown = [r for r in included if r["reference"] == "brown_figure8"]
    assert len(tested_brown) == 1
    assert tested_brown[0]["reference_low_db_pa2_hz"] == -30
    assert tested_brown[0]["window_fraction_within_bounds"] == 1
    outside = next(r for r in rows if r["reference_frequency_hz"] == 8)
    assert outside["comparison"] == "excluded"
    assert outside["window_fraction_within_bounds"] is None
    assert outside["reference_low_db_pa2_hz"] is None


def test_wrong_units_ids_and_duplicate_times_are_rejected():
    ppsd = SimpleNamespace(
        special_handling="hydrophone",
        id="XX.TEST.00.BDH",
        psd_values=[[0, 1], [1, 2]],
        times_processed=[1, 2],
        period_bin_centers=[1, 2],
    )
    row = {"Seed_id": ppsd.id, "ppsd_units": "Pa^2/Hz"}
    assert validate_pressure(ppsd, row).shape == (2, 2)
    with pytest.raises(ValueError, match="pressure PSDs"):
        validate_pressure(ppsd, row | {"ppsd_units": "m^2/s^4/Hz"})
    with pytest.raises(ValueError, match="identifier"):
        validate_pressure(ppsd, row | {"Seed_id": "XX.OTHER.00.BDH"})
    ppsd.times_processed = [1, 1]
    with pytest.raises(ValueError, match="Duplicate"):
        validate_pressure(ppsd, row)


def test_saved_pressure_integration_preserves_files_rows_and_scores(saved_pressure):
    export_dir, source, _ = saved_pressure
    before = load_master_dataset(export_dir)
    hash_before = hashlib.sha256(source.read_bytes()).hexdigest()
    result = run_pressure_test(export_dir, network="XX", station="TEST", show=False)[0]
    assert result["score"] is None
    assert result["status"] == "comparison_only"
    assert result["reference_coverage"]["brown_figure8"] > 0
    after = load_master_dataset(export_dir)
    assert after[TEST_COLUMNS[0]].tolist() == ["", "0.75"]
    assert after[
        before.columns.drop(TEST_COLUMNS[0].removesuffix("_score") + "_status")
    ].equals(before.drop(columns=TEST_COLUMNS[0].removesuffix("_score") + "_status"))
    assert after.loc[1, "pressure_expected_status"] == ""
    assert hashlib.sha256(source.read_bytes()).hexdigest() == hash_before
    book = load_workbook(export_dir / "master_inventory.xlsx")
    assert book["Manual notes"]["A1"].value == "preserve me"
    headers = {cell.value: cell.column for cell in book["Master inventory"][1]}
    assert book["Master inventory"].cell(2, headers[TEST_COLUMNS[0]]).value is None
    assert book["Master inventory"].cell(3, headers[TEST_COLUMNS[0]]).value == 0.75
    book.close()
    assert len(list((export_dir / "backups").iterdir())) == 1
    result_json = json.loads(
        (source.parent / "pressure_expected/comparison.json").read_text()
    )
    assert result_json["source_npz_sha256"] == hash_before
    for name in ("ppsd_with_limits_plot", "test_details_plot"):
        assert Path(result_json["files"][name]).is_file()


def test_score_is_explicit_and_other_tests_cannot_be_patched(saved_pressure):
    export_dir, source, ppsd = saved_pressure
    with pytest.raises(ValueError, match="explicit sensor"):
        run_pressure_test(export_dir, write_score=True, show=False)
    with pytest.raises(ValueError, match="both"):
        response_band(ppsd, source.parent / "inventory.xml", sensor_min=0.01)
    with pytest.raises(ValueError, match="Nyquist"):
        response_band(
            ppsd, source.parent / "inventory.xml", sensor_min=0.01, sensor_max=5
        )
    result = run_pressure_test(
        export_dir,
        network="XX",
        station="TEST",
        show=False,
        sensor_min=0.01,
        sensor_max=3.99,
        write_score=True,
    )[0]
    assert result["status"] == "screened"
    assert (
        float(load_master_dataset(export_dir).loc[0, TEST_COLUMNS[0]])
        == result["score"]
    )
    row = load_master_dataset(export_dir).iloc[0]
    with pytest.raises(ValueError, match="Updates may only"):
        update_pressure_test_results(
            export_dir,
            [
                {
                    **row[
                        [
                            "Seed_id",
                            "channel_start_time",
                            "channel_end_time",
                            "inventory_file",
                        ]
                    ],
                    TEST_COLUMNS[1]: 1,
                }
            ],
        )


def test_nominal_band_follows_known_highpass_response(saved_pressure):
    _, source, ppsd = saved_pressure
    snapshot = source.parent / "inventory.xml"
    inv = read_inventory(str(snapshot))
    corner, reference = 0.1, 1.0
    response = Response.from_paz(
        zeros=[0j],
        poles=[-2 * np.pi * corner + 0j],
        stage_gain=1,
        input_units="PA",
        output_units="COUNTS",
    )
    response.instrument_sensitivity.frequency = reference
    inv[0][0][0].response = response
    inv.write(str(snapshot), format="STATIONXML")
    band, _, _ = response_band(ppsd, snapshot)
    ratio = 10 ** (-3 / 10) * reference**2 / (corner**2 + reference**2)
    expected = corner * np.sqrt(ratio / (1 - ratio))
    assert band["minimum_hz"] == pytest.approx(expected, rel=0.01)
    assert band["basis"] == "nominal_response"
    assert not band["independently_verified"]


def test_reference_overlay_and_margin_plots_use_saved_diagnostics(saved_pressure):
    _, source, ppsd = saved_pressure
    metadata, brown = load_references()
    band, _, _ = response_band(ppsd, source.parent / "inventory.xml")
    band["fft_frequency_step_hz"] = ppsd.sampling_rate / ppsd.nfft
    rows, _, _ = summarize(
        np.asarray(ppsd.psd_values),
        1 / ppsd.period_bin_centers,
        1 / ppsd.period_bin_right_edges,
        1 / ppsd.period_bin_left_edges,
        band,
        metadata,
        brown,
    )
    figure = ppsd_limits_plot(ppsd, rows, band)
    axis = figure.axes[0]
    limits = {line.get_label(): line for line in axis.lines}
    assert set(limits) == {
        "Webb 1998 lower limit",
        "Webb 1998 upper limit",
        "Brown et al. 2014 lower limit",
        "Brown et al. 2014 upper limit",
    }
    assert REFERENCE_COLORS == {
        "webb_instructions": "#C62828",
        "brown_figure8": "#2E7D32",
    }
    for reference, name in [
        ("webb_instructions", "Webb 1998"),
        ("brown_figure8", "Brown et al. 2014"),
    ]:
        selected = [
            r
            for r in rows
            if r["reference"] == reference and r["reference_low_db_pa2_hz"] is not None
        ]
        for bound, label in [("low", "lower"), ("high", "upper")]:
            line = limits[f"{name} {label} limit"]
            assert line.get_color() == REFERENCE_COLORS[reference]
            np.testing.assert_allclose(
                line.get_ydata(), [r[f"reference_{bound}_db_pa2_hz"] for r in selected]
            )
    assert "Pa" in axis.get_ylabel()
    assert any("[%]" in ax.get_ylabel() for ax in figure.axes)
    plt.close(figure)
    details = make_test_details_plot(ppsd, rows)
    assert len(details.axes) == 4
    for column, reference in enumerate(REFERENCE_COLORS):
        selected = [
            r
            for r in rows
            if r["reference"] == reference and r["comparison"] != "excluded"
        ]
        points = details.axes[2 + column].collections
        actual = [tuple(collection.get_offsets()[0]) for collection in points]
        expected = [
            (r["frequency_hz"], 100 * r["window_fraction_within_bounds"])
            for r in selected
        ]
        np.testing.assert_allclose(actual, expected)
    plt.close(details)
