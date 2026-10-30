"""Check inventory epochs, column order, and lossless spreadsheet exports."""

from copy import copy

import pandas as pd
import pytest
from obspy import UTCDateTime
from obspy.core.inventory import Channel, Inventory, Network, Site, Station
from openpyxl import load_workbook

from compy_pressure.inventory import (
    DOCUMENT_COLUMNS,
    build_master_dataset,
    load_master_dataset,
    save_master_dataset,
    update_processing_results,
)


def write_inventory(path, network_code, station_code, elevation, starts):
    channels = [
        Channel(
            code="BHZ",
            location_code="00",
            latitude=10,
            longitude=20,
            elevation=elevation,
            depth=0,
            sample_rate=50,
            start_date=UTCDateTime(start),
            end_date=UTCDateTime(start) + 86400,
        )
        for start in starts
    ]
    station = Station(
        code=station_code,
        latitude=10,
        longitude=20,
        elevation=elevation,
        site=Site(name="Synthetic station"),
        channels=channels,
    )
    Inventory([Network(code=network_code, stations=[station])], source="Test").write(
        str(path), format="STATIONXML"
    )


def test_all_files_epochs_depth_and_document_order(tmp_path):
    write_inventory(
        tmp_path / "one.xml", "YV", "OBS1", -1200, ["2020-01-01", "2021-01-01"]
    )
    nested = tmp_path / "nested"
    nested.mkdir()
    write_inventory(nested / "two.XML", "ZX", "LAND1", 100, ["2020-01-01"])

    master = build_master_dataset(tmp_path)

    assert len(master) == 3
    assert list(master.columns[:10]) == DOCUMENT_COLUMNS
    assert "channel" not in master.columns
    assert "station_latitude" not in master.columns
    assert "sensor_model" not in master.columns
    assert master.loc[master.Seed_id == "YV.OBS1.00.BHZ", "Profondeur"].eq(1200).all()
    assert master.loc[master.Seed_id == "ZX.LAND1.00.BHZ", "Profondeur"].isna().all()
    assert master[DOCUMENT_COLUMNS[4:]].isna().all().all()
    assert not master.duplicate_epoch.any()


def test_missing_and_unreadable_inventories(tmp_path):
    with pytest.raises(FileNotFoundError, match="No XML files"):
        build_master_dataset(tmp_path)
    (tmp_path / "broken.xml").write_text("not XML")
    with pytest.raises(RuntimeError, match="broken.xml"):
        build_master_dataset(tmp_path)


def test_exports_preserve_ids_long_text_and_literal_strings(tmp_path):
    original = pd.DataFrame(
        {
            "Seed_id": ["YV.OBS1.00.BHZ"],
            "notes": ["=literal text"],
            "metadata": ["x" * 35000],
            "test_score": pd.Series([pd.NA], dtype="Float64"),
        }
    )
    csv_path, excel_path = save_master_dataset(original, tmp_path / "exports")
    csv = pd.read_csv(csv_path, keep_default_na=False)
    assert csv.Seed_id.iloc[0] == original.Seed_id.iloc[0]
    assert csv.metadata.iloc[0] == original.metadata.iloc[0]
    assert "metadata_part_2" not in original.columns

    workbook = load_workbook(excel_path)
    sheet = workbook["Master inventory"]
    headers = [cell.value for cell in sheet[1]]
    record = dict(zip(headers, [cell.value for cell in sheet[2]], strict=True))
    assert record["Seed_id"] == original.Seed_id.iloc[0]
    assert record["notes"] == "=literal text"
    assert sheet.cell(2, headers.index("notes") + 1).data_type == "s"
    assert (
        record["metadata"] + record["metadata_part_2"] + record["metadata_part_3"]
        == original.metadata.iloc[0]
    )
    assert record["test_score"] is None
    assert sheet.freeze_panes == "E2"
    workbook.close()


def test_processing_updates_preserve_other_epochs_scores_and_excel_features(tmp_path):
    original = pd.DataFrame(
        {
            "Seed_id": ["XX.TEST.00.BDH", "XX.TEST.00.BDH"],
            "channel_start_time": ["2020-01-01", "2021-01-01"],
            "channel_end_time": ["2020-02-01", "2021-02-01"],
            "inventory_file": ["source.xml", "source.xml"],
            "sample_rate_hz": [8, 8],
            "test_pressure_spectrum_vs_expected_noise_levels_score": [1.0, 0.0],
            "notes": ["NA", "=literal"],
        }
    )
    csv_path, excel_path = save_master_dataset(original, tmp_path)
    workbook = load_workbook(excel_path)
    workbook.create_sheet("Manual notes")["A1"] = "=1+1"
    bold_font = copy(workbook["Master inventory"]["A2"].font)
    bold_font.bold = True
    workbook["Master inventory"]["A2"].font = bold_font
    workbook.save(excel_path)
    workbook.close()
    update = original.iloc[0][
        [
            "Seed_id",
            "channel_start_time",
            "channel_end_time",
            "inventory_file",
        ]
    ].to_dict() | {
        "ppsd_status": "complete",
        "ppsd_segments": 7,
        "ppsd_units": "Pa^2/Hz",
    }
    update_processing_results(tmp_path, [update])
    result = load_master_dataset(tmp_path)
    assert result["ppsd_status"].tolist() == ["complete", "not_processed"]
    assert result["ppsd_segments"].tolist() == ["7", ""]
    assert result["notes"].tolist() == ["NA", "=literal"]
    assert result[original.columns].equals(
        pd.read_csv(
            next((tmp_path / "backups").glob("*/master_inventory.csv")),
            dtype=str,
            keep_default_na=False,
        )
    )
    workbook = load_workbook(excel_path)
    sheet = workbook["Master inventory"]
    headers = {cell.value: cell.column for cell in sheet[1]}
    assert sheet.cell(2, headers["ppsd_segments"]).value == 7
    assert sheet.cell(3, headers["ppsd_segments"]).value is None
    assert (
        sheet.cell(
            2, headers["test_pressure_spectrum_vs_expected_noise_levels_score"]
        ).value
        == 1
    )
    assert sheet["A2"].font.bold
    assert sheet.freeze_panes == "E2"
    assert workbook["Manual notes"]["A1"].value == "=1+1"
    workbook.close()
    assert csv_path.exists()


def test_processing_keys_remain_even_when_metadata_columns_are_identical(tmp_path):
    write_inventory(tmp_path / "one.xml", "XX", "OBS", -1000, ["2020-01-01"])
    master = build_master_dataset(tmp_path)
    assert {
        "channel_start_time",
        "channel_end_time",
        "sample_rate_hz",
        "inventory_file",
    }.issubset(master.columns)


def test_csv_excel_row_disagreement_refuses_both_updates(tmp_path):
    import hashlib

    original = pd.DataFrame(
        {
            "Seed_id": ["XX.TEST.00.BDH", "XX.TEST.00.BHZ"],
            "channel_start_time": ["2020-01-01", "2020-01-01"],
            "channel_end_time": ["2020-02-01", "2020-02-01"],
            "inventory_file": ["source.xml", "source.xml"],
            "sample_rate_hz": [8, 8],
        }
    )
    _, excel_path = save_master_dataset(original, tmp_path)
    workbook = load_workbook(excel_path)
    workbook["Master inventory"].delete_rows(3)
    workbook.save(excel_path)
    workbook.close()
    hashes = {
        p: hashlib.sha256(p.read_bytes()).hexdigest()
        for p in tmp_path.glob("master_inventory.*")
    }
    key = original.iloc[1][
        ["Seed_id", "channel_start_time", "channel_end_time", "inventory_file"]
    ].to_dict()
    with pytest.raises(ValueError, match="disagree"):
        update_processing_results(tmp_path, [key | {"ppsd_status": "complete"}])
    assert hashes == {p: hashlib.sha256(p.read_bytes()).hexdigest() for p in hashes}
    assert not (tmp_path / "backups").exists()


@pytest.mark.parametrize("existing", ["both", "csv_only"])
def test_full_export_refuses_to_replace_an_existing_master(tmp_path, existing):
    import hashlib

    scored = pd.DataFrame(
        {
            "Seed_id": ["XX.TEST.00.BDH"],
            "channel_start_time": ["2020-01-01"],
            "channel_end_time": ["2020-02-01"],
            "inventory_file": ["source.xml"],
            "test_pressure_spectrum_vs_expected_noise_levels_score": [1.0],
        }
    )
    _, excel_path = save_master_dataset(scored, tmp_path)
    if existing == "csv_only":
        excel_path.unlink()
    hashes = {
        p: hashlib.sha256(p.read_bytes()).hexdigest()
        for p in tmp_path.glob("master_inventory.*")
    }
    fresh = scored.drop(columns="test_pressure_spectrum_vs_expected_noise_levels_score")
    with pytest.raises(FileExistsError, match="network inventory"):
        save_master_dataset(fresh, tmp_path)
    assert hashes == {
        p: hashlib.sha256(p.read_bytes()).hexdigest()
        for p in tmp_path.glob("master_inventory.*")
    }


def test_master_update_waits_for_another_writer(tmp_path):
    import fcntl
    import threading

    original = pd.DataFrame(
        {
            "Seed_id": ["XX.TEST.00.BDH", "XX.TEST.00.BHZ"],
            "channel_start_time": ["2020-01-01", "2020-01-01"],
            "channel_end_time": ["2020-02-01", "2020-02-01"],
            "inventory_file": ["source.xml", "source.xml"],
            "sample_rate_hz": [8, 8],
        }
    )
    save_master_dataset(original, tmp_path)
    key = original.iloc[0][
        ["Seed_id", "channel_start_time", "channel_end_time", "inventory_file"]
    ].to_dict()
    writer = threading.Thread(
        target=update_processing_results,
        args=(tmp_path, [key | {"ppsd_status": "complete"}]),
        daemon=True,
    )
    with (tmp_path / ".master_inventory.lock").open("a+") as other_process:
        fcntl.flock(other_process, fcntl.LOCK_EX)
        writer.start()
        writer.join(timeout=1)
        # A concurrent read-modify-write would otherwise drop the holder's fields.
        assert writer.is_alive()
        assert not (tmp_path / "backups").exists()
    writer.join(timeout=60)
    assert not writer.is_alive()
    assert load_master_dataset(tmp_path).ppsd_status.tolist() == [
        "complete",
        "not_processed",
    ]


def test_fresh_export_obeys_shared_master_lock(tmp_path):
    import fcntl
    import threading

    errors = []
    table = pd.DataFrame({"Seed_id": ["XX.TEST..BDH"]})

    def export():
        try:
            save_master_dataset(table, tmp_path)
        except Exception as exc:
            errors.append(exc)

    writer = threading.Thread(target=export, daemon=True)
    with (tmp_path / ".master_inventory.lock").open("a+") as holder:
        fcntl.flock(holder, fcntl.LOCK_EX)
        writer.start()
        writer.join(timeout=1)
        assert writer.is_alive(), "Fresh export bypassed the shared master lock"
        # Another cooperating writer creates the master while holding the lock.
        protected = tmp_path / "master_inventory.csv"
        protected.write_text("Seed_id,manual_score\nXX.OTHER..BDH,1\n")
        before = protected.read_bytes()
    writer.join(timeout=10)
    assert not writer.is_alive()
    assert len(errors) == 1 and isinstance(errors[0], FileExistsError)
    assert protected.read_bytes() == before
    assert not (tmp_path / "master_inventory.xlsx").exists()
