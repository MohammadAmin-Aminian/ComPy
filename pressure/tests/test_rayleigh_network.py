"""Network isolation, restricted master updates, epoch binding and restart."""

import hashlib
import json
from pathlib import Path

import numpy as np
import pandas as pd
import pytest
from openpyxl import load_workbook
from test_rayleigh_calibration import synthetic_station as synthetic_station

from compy_pressure import network, rayleigh_network
from compy_pressure.inventory import (
    RESULT_COLUMNS,
    load_master_dataset,
    update_rayleigh_results,
)
from compy_pressure.rayleigh_calibration import select_calibration_channels


def options(catalog, dispersion):
    return dict(
        catalog=catalog,
        dispersion_file=dispersion,
        pressure_sign=1,
        band=[0.02, 0.06],
        max_events=1,
    )


def test_real_batch_preserves_master_scores_sheets_and_resumes(
    synthetic_station, tmp_path, monkeypatch
):
    data, exports, catalog, dispersion = synthetic_station
    book = load_workbook(exports / "master_inventory.xlsx")
    book.create_sheet("Manual notes")["A1"] = "preserve me"
    book.save(exports / "master_inventory.xlsx")
    book.close()
    before = load_master_dataset(exports)
    raw = {p: hashlib.sha256(p.read_bytes()).hexdigest() for p in data.iterdir()}
    files = network.scan_files(data)
    results = tmp_path / "network_results"
    report = rayleigh_network.run_network(
        exports, results, files, options=options(catalog, dispersion)
    )
    assert report["status_counts"] == {"unresolved": 1}
    entry = next(iter(report["entries"].values()))
    saved = json.loads(Path(entry["update"]["rayleigh_result_file"]).read_text())
    assert saved["candidate_events"] == 1
    assert (
        saved["implementation_status"] == "working_method_with_provisional_assumptions"
    )
    assert (
        "earth_model_needs_station_validation"
        in entry["update"]["rayleigh_review_flags"]
    )
    after = load_master_dataset(exports)
    pd.testing.assert_frame_equal(before, after[before.columns])
    book = load_workbook(exports / "master_inventory.xlsx")
    assert book["Manual notes"]["A1"].value == "preserve me"
    book.close()
    assert raw == {
        p: hashlib.sha256(p.read_bytes()).hexdigest() for p in data.iterdir()
    }
    original = rayleigh_network.run_calibration
    with monkeypatch.context() as patch:
        patch.setattr(
            rayleigh_network,
            "run_calibration",
            lambda **kwargs: pytest.fail("cached result was recomputed"),
        )
        resumed = rayleigh_network.run_network(
            exports, results, files, options=options(catalog, dispersion)
        )
    assert (
        next(iter(resumed["entries"].values()))["fingerprint"] == entry["fingerprint"]
    )
    next(p for p in data.iterdir() if "BHZ" in p.name).touch()
    monkeypatch.setattr(rayleigh_network, "run_calibration", original)
    changed = rayleigh_network.run_network(
        exports, results, files, options=options(catalog, dispersion)
    )
    assert (
        next(iter(changed["entries"].values()))["fingerprint"] != entry["fingerprint"]
    )
    with pytest.raises(ValueError, match="rayleigh"):
        update_rayleigh_results(
            exports,
            [{k: before.iloc[0][k] for k in network.KEYS} | {RESULT_COLUMNS[0]: 1}],
        )


def test_no_tilt_network_and_separate_workflow_view(
    synthetic_station, tmp_path, monkeypatch
):
    data, exports, catalog, dispersion = synthetic_station
    master = load_master_dataset(exports)
    master = master[~master.Seed_id.str.endswith(("BH1", "BH2"))]
    master.to_csv(exports / "master_inventory.csv", index=False)
    for path in data.iterdir():
        if path.name.split(".")[3] in {"BH1", "BH2"}:
            path.unlink()
    before = {
        p: hashlib.sha256(p.read_bytes()).hexdigest()
        for p in exports.glob("master_inventory.*")
    }
    monkeypatch.setattr(
        rayleigh_network,
        "require_clean_rotator",
        lambda: pytest.fail("Network required TiSKitPy without tilt"),
    )
    files = network.scan_files(data)
    results = tmp_path / "no_tilt_results"
    report = rayleigh_network.run_network(
        exports,
        results,
        files,
        update_master=False,
        options=options(catalog, dispersion) | {"tilt_method": "none"},
    )
    entry = next(iter(report["entries"].values()))
    assert entry["status"] == "unresolved"  # Single event, not a metadata failure.
    saved = json.loads(Path(entry["update"]["rayleigh_result_file"]).read_text())
    assert len(saved["epoch_keys"]) == 2
    assert saved["events"][0]["tilt"]["method"] == "none"
    assert not report["master_updated"]
    paths = network.view_plots(
        exports,
        files,
        "rayleigh",
        show=False,
        workflow_path=results / "_workflow/rayleigh.json",
    )
    assert paths == [Path(entry["update"]["rayleigh_plot_file"])]
    assert before == {p: hashlib.sha256(p.read_bytes()).hexdigest() for p in before}
    # Explicit workflow scope must not fall back to historical master plots.
    empty = results / "empty.json"
    empty.write_text(json.dumps({"entries": {}}))
    with pytest.raises(ValueError, match="No saved rayleigh plots"):
        network.view_plots(exports, files, "rayleigh", show=False, workflow_path=empty)


def test_missing_component_is_recorded_without_master_changes(
    synthetic_station, tmp_path
):
    data, exports, catalog, dispersion = synthetic_station
    master = load_master_dataset(exports)
    master = master[~master.Seed_id.str.endswith("BH2")]
    master.to_csv(exports / "master_inventory.csv", index=False)
    hashes = {
        p: hashlib.sha256(p.read_bytes()).hexdigest()
        for p in exports.glob("master_inventory.*")
    }
    report = rayleigh_network.run_network(
        exports,
        tmp_path / "results",
        network.scan_files(data),
        update_master=False,
        options=options(catalog, dispersion),
    )
    assert report["status_counts"] == {"unresolved_metadata": 1}
    entry = next(iter(report["entries"].values()))
    assert "BH2" in entry["error"]
    assert entry["update"]["rayleigh_pressure_multiplier"] == ""
    assert hashes == {p: hashlib.sha256(p.read_bytes()).hexdigest() for p in hashes}


def test_unusable_reference_band_stops_before_raw_reads(
    synthetic_station, tmp_path, monkeypatch
):
    data, exports, catalog, dispersion = synthetic_station
    # Cover the requested band so this isolates the resonance preflight, rather
    # than failing the separate prohibition on dispersion extrapolation.
    np.savetxt(
        dispersion,
        [[0.015, 4000, 3800], [0.45, 4000, 3800]],
        delimiter=",",
        header="frequency_hz,phase_velocity_m_s,group_velocity_m_s",
        comments="",
    )
    monkeypatch.setattr(
        rayleigh_network,
        "index_waveform_files",
        lambda *args: pytest.fail("unusable band caused raw read"),
    )
    report = rayleigh_network.run_network(
        exports,
        tmp_path / "results",
        network.scan_files(data),
        update_master=False,
        options=options(catalog, dispersion) | dict(band=[0.03, 0.4]),
    )
    assert report["status_counts"] == {"reference_band_unusable": 1}
    assert (
        next(iter(report["entries"].values()))["update"]["rayleigh_pressure_multiplier"]
        == ""
    )


def test_pressure_epoch_binding_refuses_ambiguous_station(synthetic_station):
    _, exports, _, _ = synthetic_station
    master = load_master_dataset(exports)
    row = master[master.input_units.eq("PA")].iloc[0]
    other = row.copy()
    other.channel_start_time = "2010-01-01T00:00:00.000000Z"
    expanded = pd.concat([master, other.to_frame().T], ignore_index=True)
    with pytest.raises(ValueError, match="one pressure epoch"):
        select_calibration_channels(expanded, "XX", "TEST")
    rows, _ = select_calibration_channels(expanded, "XX", "TEST", row)
    assert rows[0].channel_start_time == row.channel_start_time


def test_active_other_workflow_blocks_master_writes(synthetic_station, tmp_path):
    data, exports, catalog, dispersion = synthetic_station
    results = tmp_path / "results"
    (results / "_workflow").mkdir(parents=True)
    (results / "_workflow/paired.json").write_text(
        json.dumps(dict(selected_epoch_keys=["running"], entries={}))
    )
    with pytest.raises(ValueError, match="Unfinished paired"):
        rayleigh_network.run_network(
            exports,
            results,
            network.scan_files(data),
            options=options(catalog, dispersion),
        )


@pytest.mark.parametrize("units", ["Pa", "PASCAL", " Pascals "])
def test_rayleigh_accepts_supported_pressure_unit_aliases(
    synthetic_station, units, tmp_path, monkeypatch
):
    data, exports, catalog, dispersion = synthetic_station
    master = load_master_dataset(exports)
    master.loc[master.input_units.eq("PA"), "input_units"] = units
    rows, _ = select_calibration_channels(master, "XX", "TEST", tilt_method="none")
    assert rows[0].Seed_id.endswith("BDH")
    master.to_csv(exports / "master_inventory.csv", index=False)

    def reached_calibration(**kwargs):
        raise ValueError("Supported pressure alias reached calibration")

    monkeypatch.setattr(rayleigh_network, "run_calibration", reached_calibration)
    report = rayleigh_network.run_network(
        exports,
        tmp_path / "results",
        network.scan_files(data),
        update_master=False,
        options=options(catalog, dispersion) | {"tilt_method": "none"},
    )
    assert report["status_counts"] == {"unresolved_metadata": 1}
    assert "pressure alias" in next(iter(report["entries"].values()))["error"]


@pytest.mark.parametrize("coordinate", ["Lat", "Lon"])
def test_changed_event_geometry_invalidates_cache(synthetic_station, coordinate):
    data, exports, catalog, dispersion = synthetic_station
    master = load_master_dataset(exports)
    row = master[master.input_units.eq("PA")].iloc[0].copy()
    files = network.scan_files(data)
    settings = options(str(catalog), str(dispersion))
    stamp = rayleigh_network.dependency_stamp(row, master, files, settings)
    row[coordinate] = str(float(row[coordinate]) + 1)
    assert rayleigh_network.dependency_stamp(row, master, files, settings) != stamp


def test_network_preflight_uses_explicit_depth(
    synthetic_station, tmp_path, monkeypatch
):
    data, exports, catalog, dispersion = synthetic_station
    np.savetxt(
        dispersion,
        [[0.015, 4000, 3800], [0.45, 4000, 3800]],
        delimiter=",",
        header="frequency_hz,phase_velocity_m_s,group_velocity_m_s",
        comments="",
    )
    calls = []

    def record_call(**kwargs):
        calls.append(kwargs["depth"])
        raise ValueError("Reached calibration with reviewed depth")

    monkeypatch.setattr(rayleigh_network, "run_calibration", record_call)
    report = rayleigh_network.run_network(
        exports,
        tmp_path / "results",
        network.scan_files(data),
        update_master=False,
        options=options(catalog, dispersion) | dict(band=[0.03, 0.4], depth=500),
    )
    assert calls == [500]
    assert report["status_counts"] == {"unresolved_metadata": 1}
    assert "reviewed depth" in next(iter(report["entries"].values()))["error"]
