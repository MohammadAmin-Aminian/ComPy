"""Independent normalization/transfer checks and real MiniSEED preservation."""

import hashlib
import json
from pathlib import Path

import numpy as np
import pytest
from obspy import Trace, UTCDateTime
from obspy.core.inventory import Channel, Inventory, Network, Site, Station
from obspy.core.inventory.response import Response
from openpyxl import load_workbook
from scipy.signal import periodogram

from compy_pressure import network, paired_spectra
from compy_pressure.inventory import (
    RESULT_COLUMNS,
    load_master_dataset,
    synchronize_inventory,
    update_paired_results,
    update_processing_results,
)


@pytest.mark.parametrize("n", [4095, 4096])
def test_density_against_independent_scipy_and_parseval(n):
    rng = np.random.default_rng(18)
    p = rng.normal(size=n)
    rp = np.full(n // 2 + 1, 4 + 0j)
    rz = np.full(n // 2 + 1, 8 + 0j)
    pp, zz, zp, valid = paired_spectra.spectral_terms(p * 4, p * 24, 8, rp, rz)
    f, expected = periodogram(p, fs=8, window="hann", detrend="constant")
    np.testing.assert_allclose(pp[1:], expected[1:], rtol=1e-12, atol=1e-16)
    np.testing.assert_allclose(zz[1:], 9 * expected[1:], rtol=1e-12, atol=1e-16)
    np.testing.assert_allclose(zp[1:], 3 * expected[1:], rtol=1e-12, atol=1e-16)
    taper = 0.5 - 0.5 * np.cos(2 * np.pi * np.arange(n) / n)
    energy = np.sum(((p - p.mean()) * taper) ** 2) / np.sum(taper**2)
    assert (pp.sum() + expected[0]) * (f[1] - f[0]) == pytest.approx(energy)
    assert not valid[0]


def test_known_complex_response_polarity_and_average_before_coherence():
    rng = np.random.default_rng(9)
    terms = []
    for _ in range(30):
        p = rng.normal(size=1024)
        rp = np.full(513, 4 * np.exp(0.3j))
        rz = np.full(513, 8 * np.exp(-0.5j))
        terms.append(paired_spectra.spectral_terms(p * 4, p * 24, 4, rp, rz, -1))
    pp, zz, zp, h, coherence, valid = paired_spectra.averaged_products(
        *[np.array([t[i] for t in terms]) for i in range(3)]
    )
    np.testing.assert_allclose(h[valid], -3 * np.exp(0.8j), rtol=1e-12)
    np.testing.assert_allclose(coherence[valid], 1, atol=1e-12)
    np.testing.assert_allclose(zp[valid], h[valid] * pp[valid])
    assert np.all(zz[valid] > 0)
    with pytest.raises(ValueError, match="two"):
        paired_spectra.averaged_products([terms[0][0]], [terms[0][1]], [terms[0][2]])


def test_saved_response_can_be_standalone_without_ppsd(tmp_path):
    snapshot = tmp_path / "inventory.xml"
    snapshot.write_text("preserved snapshot")
    row = {"rayleigh_response_snapshot_file": str(snapshot)}
    assert paired_spectra.response_snapshot_path(row) == snapshot
    assert paired_spectra.has_saved_response(row)
    assert not paired_spectra.has_saved_response({})


def test_independent_noise_and_known_delay_phase():
    rng = np.random.default_rng(81)
    n, rate, delay = 2048, 8, 0.25
    f = np.fft.rfftfreq(n, 1 / rate)
    independent, delayed = [], []
    for _ in range(80):
        p, z = rng.normal(size=(2, n))
        independent.append(
            paired_spectra.spectral_terms(p, z, rate, np.ones(len(f)), np.ones(len(f)))
        )
        # A known response delay should produce the opposite correction phase.
        delayed.append(
            paired_spectra.spectral_terms(
                p, p, rate, np.ones(len(f)), np.exp(-2j * np.pi * f * delay)
            )
        )
    products = paired_spectra.averaged_products(
        *[np.array([t[i] for t in independent]) for i in range(3)]
    )
    assert np.mean(products[4][1:]) < 0.025
    result = paired_spectra.averaged_products(
        *[np.array([t[i] for t in delayed]) for i in range(3)]
    )
    np.testing.assert_allclose(
        result[3][1:], np.exp(2j * np.pi * f[1:] * delay), atol=1e-12
    )


@pytest.fixture
def paired_data(tmp_path):
    start = UTCDateTime("2020-01-02T00:00:00.100000")
    rate = 2
    data, inventories, exports, results = [
        tmp_path / n for n in ("data", "inventories", "exports", "results")
    ]
    data.mkdir()
    inventories.mkdir()
    channels = []
    rng = np.random.default_rng(28)
    physical = rng.normal(size=2 * 86400 + 8).astype(np.float64)
    for code, unit, gain, dip in [("BDH", "PA", 4, 0), ("BHZ", "M", 8, -90)]:
        response = Response.from_paz(
            zeros=[], poles=[], stage_gain=gain, input_units=unit, output_units="COUNTS"
        )
        response.instrument_sensitivity.frequency = 0.2
        channels.append(
            Channel(
                code=code,
                location_code="00",
                latitude=1,
                longitude=2,
                elevation=-1000,
                depth=0,
                sample_rate=rate,
                dip=dip,
                start_date=start - 86400,
                end_date=start + 2 * 86400,
                response=response,
            )
        )
        trace = Trace(
            physical * gain * (3 if code == "BHZ" else 1),
            header={
                "network": "XX",
                "station": "PAIR",
                "location": "00",
                "channel": code,
                "sampling_rate": rate,
                "starttime": start,
            },
        )
        trace.write(str(data / f"XX.PAIR.00.{code}.D.2020.002"), format="MSEED")
    inv = Inventory(
        [
            Network(
                "XX",
                stations=[
                    Station(
                        "PAIR",
                        latitude=1,
                        longitude=2,
                        elevation=-1000,
                        site=Site(name="Pair"),
                        channels=channels,
                    )
                ],
            )
        ],
        "Test",
    )
    inv.write(str(inventories / "XX.xml"), format="STATIONXML")
    synchronize_inventory(inventories, exports)
    master = load_master_dataset(exports)
    patches = []
    for i, row in master.iterrows():
        folder = results / "XX" / "PAIR" / str(i)
        folder.mkdir(parents=True)
        inv.select(channel=row.Seed_id.split(".")[-1]).write(
            str(folder / "inventory.xml"), format="STATIONXML"
        )
        (folder / "ppsd.npz").write_bytes(b"preserve this saved product")
        patches.append(
            {n: row[n] for n in paired_spectra.KEYS}
            | {"ppsd_status": "complete", "ppsd_npz_file": str(folder / "ppsd.npz")}
        )
    update_processing_results(exports, patches)
    book = load_workbook(exports / "master_inventory.xlsx")
    book.create_sheet("Manual notes")["A1"] = "keep"
    book.save(exports / "master_inventory.xlsx")
    book.close()
    return data, exports, results, start


def test_real_pair_master_preservation_network_cache_and_dependencies(
    paired_data, monkeypatch
):
    data, exports, results, start = paired_data
    before = load_master_dataset(exports)
    raw = {p: hashlib.sha256(p.read_bytes()).hexdigest() for p in data.iterdir()}
    options = dict(start=str(start), days=1, segment_seconds=3600)
    report = network.run_stage(
        "paired", exports, results, network.scan_files(data), options=options
    )
    assert report["status_counts"] == {"comparison_only": 1}
    entry = next(iter(report["entries"].values()))
    summary = json.loads(
        Path(
            next(p for p in entry["files"] if p.endswith("comparison.json"))
        ).read_text()
    )
    assert summary["retained_segments"] == 24
    assert (
        paired_spectra.view_saved(summary["files"]["result"], show=False)
        == summary["files"]["plot"]
    )
    with np.load(summary["files"]["spectra"]) as saved:
        np.testing.assert_allclose(saved["transfer_dp"][saved["valid"]], 3, atol=1e-8)
        np.testing.assert_allclose(
            saved["coherence_squared"][saved["valid"]], 1, atol=1e-10
        )
    after = load_master_dataset(exports)
    assert before.equals(after[before.columns])
    assert after[RESULT_COLUMNS].eq("").all().all()
    book = load_workbook(exports / "master_inventory.xlsx")
    assert book["Manual notes"]["A1"].value == "keep"
    book.close()
    assert raw == {
        p: hashlib.sha256(p.read_bytes()).hexdigest() for p in data.iterdir()
    }
    original = network.diagnostic_job
    monkeypatch.setattr(
        network, "diagnostic_job", lambda *a, **k: pytest.fail("cache missed")
    )
    cached = network.run_stage(
        "paired", exports, results, network.scan_files(data), options=options
    )
    assert next(iter(cached["entries"].values()))["files"] == entry["files"]
    monkeypatch.setattr(network, "diagnostic_job", original)
    vertical = next(p for p in data.iterdir() if "BHZ" in p.name)
    vertical.touch()
    rerun = network.run_stage(
        "paired", exports, results, network.scan_files(data), options=options
    )
    assert next(iter(rerun["entries"].values()))["fingerprint"] != entry["fingerprint"]


def test_unaligned_grids_and_short_coverage_are_saved_unresolved(paired_data):
    from obspy import read

    data, exports, _, start = paired_data
    vertical = next(p for p in data.iterdir() if "BHZ" in p.name)
    stream = read(str(vertical))
    stream[0].stats.starttime += 0.1
    stream.write(str(vertical), format="MSEED")
    original = {
        p: hashlib.sha256(p.read_bytes()).hexdigest()
        for p in exports.glob("master_inventory.*")
    }
    summary = paired_spectra.run_paired_spectra(
        exports, data_dir=data, network="XX", station="PAIR", start=start + 1, days=1
    )
    assert summary["status"] == "no_complete_pairs"
    assert summary["retained_segments"] == 0
    summary = paired_spectra.run_paired_spectra(
        exports, data_dir=data, network="XX", station="PAIR", days=7
    )
    assert summary["status"] == "insufficient_common_coverage"
    assert original == {
        p: hashlib.sha256(p.read_bytes()).hexdigest()
        for p in exports.glob("master_inventory.*")
    }


def test_restricted_update_and_unfinished_workflow_guard(paired_data):
    _, exports, results, _ = paired_data
    row = load_master_dataset(exports).iloc[0]
    with pytest.raises(ValueError, match="paired"):
        update_paired_results(
            exports, [{n: row[n] for n in paired_spectra.KEYS} | {RESULT_COLUMNS[0]: 1}]
        )
    folder = results / "_workflow"
    folder.mkdir()
    (folder / "pilot.json").write_text(
        json.dumps({"selected_epoch_keys": ["a"], "status_counts": {}})
    )
    with pytest.raises(ValueError, match="Unfinished pilot"):
        paired_spectra.assert_workflows_finished(results)


def test_pair_metadata_units_orientation_ambiguity_and_rate(paired_data):
    from obspy import read_inventory

    _, exports, _, _ = paired_data
    master = load_master_dataset(exports)
    pressure = master[master.input_units.eq("PA")].iloc[0]
    with pytest.raises(paired_spectra.PairingError, match="found 0"):
        paired_spectra.select_pair(pressure, master[master.input_units.eq("PA")])
    duplicate = master[master.input_units.eq("M")].copy()
    duplicate.inventory_file = str(duplicate.iloc[0].inventory_file) + ".copy.xml"
    Path(duplicate.iloc[0].inventory_file).write_bytes(
        Path(pressure.inventory_file).read_bytes()
    )
    import pandas as pd

    with pytest.raises(paired_spectra.PairingError, match="found 2"):
        paired_spectra.select_pair(pressure, pd.concat([master, duplicate]))
    inv = read_inventory(pressure.inventory_file)
    inv[0][0][1].dip = 0
    inv.write(pressure.inventory_file, format="STATIONXML")
    with pytest.raises(paired_spectra.PairingError, match="found 0"):
        paired_spectra.select_pair(pressure, master)
    inv[0][0][1].dip = -90
    inv[0][0][1].sample_rate = 3
    inv.write(pressure.inventory_file, format="STATIONXML")
    with pytest.raises(ValueError, match="Unequal sample rates"):
        paired_spectra.select_pair(pressure, master)
    inv[0][0][1].sample_rate = 2
    inv[0][0][1].response.instrument_sensitivity.input_units = "PA"
    inv.write(pressure.inventory_file, format="STATIONXML")
    with pytest.raises(paired_spectra.PairingError, match="found 0"):
        paired_spectra.select_pair(pressure, master)


def test_flat_and_gapped_segments_are_excluded_without_fill(paired_data):
    from obspy import Stream, read

    data, exports, _, start = paired_data
    vertical = next(p for p in data.iterdir() if "BHZ" in p.name)
    trace = read(str(vertical))[0]
    first = trace.slice(start, start + 3600 - 0.5).copy()
    first.data[:] = 5
    rest = trace.slice(start + 7200).copy()
    Stream([first, rest]).write(str(vertical), format="MSEED")
    summary = paired_spectra.run_paired_spectra(
        exports, data_dir=data, network="XX", station="PAIR", start=start, days=1
    )
    assert summary["status"] == "comparison_only"
    assert summary["retained_segments"] == 22
    import csv

    decisions = list(csv.DictReader(Path(summary["files"]["windows"]).open()))
    assert [d["decision"] for d in decisions[:2]] == [
        "missing_gapped_nonfinite_or_flat"
    ] * 2


def test_invalid_response_bins_are_explicitly_masked():
    p = np.arange(16, dtype=float)
    response = np.ones(9, complex)
    response[2] = 0
    response[3] = np.nan
    pp, zz, zp, valid = paired_spectra.spectral_terms(p, p, 4, response, np.ones(9))
    assert not valid[[0, 2, 3]].any()
    products = paired_spectra.averaged_products([pp, pp], [zz, zz], [zp, zp])
    assert not products[-1][[0, 2, 3]].any()
    assert np.isfinite(products[3]).all()


def test_velocity_sensor_displacement_transfer_from_time_domain_integral():
    rate, n = 4, 1024
    f = np.fft.rfftfreq(n, 1 / rate)
    rz = Response.from_paz(
        zeros=[], poles=[], stage_gain=8, input_units="M/S", output_units="COUNTS"
    ).get_evalresp_response_for_frequencies(f, output="DISP")
    rp = np.full(f.size, 4 + 0j)
    # Whole cycles per segment at every other bin isolate the unit conventions.
    f0 = f[14:116:2]
    t = np.arange(4 * n) / rate
    phase = np.random.default_rng(12).uniform(0, 2 * np.pi, f0.size)
    displacement = np.cos(2 * np.pi * f0 * t[:, None] + phase).sum(axis=1)
    # A velocity sensor records the analytic derivative; DISP must integrate it.
    velocity = (-2 * np.pi * f0 * np.sin(2 * np.pi * f0 * t[:, None] + phase)).sum(
        axis=1
    )
    compliance = 2e-7  # Displacement per pressure, m/Pa
    terms = [
        paired_spectra.spectral_terms(
            4 * displacement[i : i + n] / compliance,
            8 * velocity[i : i + n],
            rate,
            rp,
            rz,
        )
        for i in range(0, 4 * n, n)
    ]
    _, _, _, h, coherence, valid = paired_spectra.averaged_products(
        *[np.array([term[i] for term in terms]) for i in range(3)]
    )
    at = np.isin(f, f0)
    assert valid[at].all()
    np.testing.assert_allclose(h[at], compliance, rtol=1e-9)
    np.testing.assert_allclose(coherence[at], 1, atol=1e-9)
