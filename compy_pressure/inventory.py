"""Build the master channel table from all StationXML files in a folder."""

# %%
import argparse
import shutil
from contextlib import contextmanager
from copy import copy
from datetime import datetime, timezone
from pathlib import Path
from tempfile import TemporaryDirectory

import pandas as pd
from obspy import read_inventory
from openpyxl import load_workbook
from openpyxl.styles import Alignment

from compy_pressure.common import PROJECT_DIRECTORY

try:
    import fcntl
except ImportError:  # Windows
    import msvcrt

    fcntl = None

XML_DIRECTORY = Path.home() / "Desktop/Data_Bruit_FM/INVENTORIES"
OUTPUT_DIRECTORY = PROJECT_DIRECTORY / "exports"

# The document lists six checks under "Calibration tests". Separate the three
# spectral-level tests from the three pressure/vertical-motion calibration checks.
# These scores assess calibration; they are not instrument correction factors.
# 1 = pass, 0 = fail, 0 < score < 1 = partial pass; blank = no result.
TEST_COLUMNS = [
    "test_pressure_spectrum_vs_expected_noise_levels_score",
    "test_spectrum_vs_nearby_stations_at_similar_depth_score",
    "test_pressure_spectrum_vs_predicted_infragravity_levels_score",
]
CALIBRATION_CHECK_COLUMNS = [
    "calibration_check_pressure_vertical_rayleigh_waves_score",
    "calibration_check_pressure_vertical_microseisms_score",
    "calibration_check_pressure_vertical_local_p_waves_score",
]
RESULT_COLUMNS = TEST_COLUMNS + CALIBRATION_CHECK_COLUMNS
DOCUMENT_COLUMNS = ["Seed_id", "Lat", "Lon", "Profondeur"] + RESULT_COLUMNS
EPOCH_COLUMNS = ["Seed_id", "channel_start_time", "channel_end_time"]
# These fields are needed to retrieve the full response and update the right epoch.
PROCESSING_METADATA_COLUMNS = [
    "channel_start_time",
    "channel_end_time",
    "sample_rate_hz",
    "inventory_file",
    "input_units",
    "output_units",
    "response_stage_count",
    "duplicate_epoch",
]


def iso_date(value):
    return str(value) if value is not None else None


def build_master_dataset(directory):
    """Read every XML; retain one row per source channel epoch, without filtering."""
    xml_files = sorted(
        path
        for path in Path(directory).rglob("*")
        if path.is_file() and path.suffix.lower() == ".xml"
    )
    if not xml_files:
        raise FileNotFoundError(f"No XML files found in {directory}")

    rows = []
    for xml_file in xml_files:
        try:
            inventory = read_inventory(str(xml_file), format="STATIONXML")
        except Exception as exc:
            raise RuntimeError(f"Could not read StationXML: {xml_file}") from exc

        for network in inventory:
            for station in network:
                for channel in station:
                    response = channel.response
                    sensitivity = response.instrument_sensitivity if response else None
                    sensor = channel.sensor
                    rows.append(
                        {
                            "Seed_id": (
                                f"{network.code}.{station.code}."
                                f"{channel.location_code}.{channel.code}"
                            ),
                            "network": network.code,
                            "station": station.code,
                            "location": channel.location_code,
                            "channel": channel.code,
                            "station_name": station.site.name if station.site else None,
                            "network_start_time": iso_date(network.start_date),
                            "network_end_time": iso_date(network.end_date),
                            "station_start_time": iso_date(station.start_date),
                            "station_end_time": iso_date(station.end_date),
                            "channel_start_time": iso_date(channel.start_date),
                            "channel_end_time": iso_date(channel.end_date),
                            "Lat": channel.latitude,
                            "Lon": channel.longitude,
                            # Estimated surface depth below sea level, not bathymetry.
                            "Profondeur": (
                                -station.elevation
                                if station.elevation is not None
                                and station.elevation < 0
                                else None
                            ),
                            "sample_rate_hz": channel.sample_rate,
                            "azimuth_deg": channel.azimuth,
                            "dip_deg": channel.dip,
                            "sensor_type": sensor.type if sensor else None,
                            "sensor_description": sensor.description
                            if sensor
                            else None,
                            "sensor_manufacturer": sensor.manufacturer
                            if sensor
                            else None,
                            "sensor_model": sensor.model if sensor else None,
                            "sensor_serial_number": sensor.serial_number
                            if sensor
                            else None,
                            "sensor_calibration_dates": (
                                "; ".join(
                                    str(date) for date in sensor.calibration_dates
                                )
                                if sensor and sensor.calibration_dates
                                else None
                            ),
                            "sensitivity": sensitivity.value if sensitivity else None,
                            "sensitivity_frequency_hz": (
                                sensitivity.frequency if sensitivity else None
                            ),
                            "input_units": sensitivity.input_units
                            if sensitivity
                            else None,
                            "output_units": sensitivity.output_units
                            if sensitivity
                            else None,
                            "response_stage_count": (
                                len(response.response_stages) if response else 0
                            ),
                            "inventory_file": str(xml_file),
                        }
                    )

    if not rows:
        raise ValueError("The inventories contain no channel records")
    master = pd.DataFrame(rows)
    for column in RESULT_COLUMNS:
        master[column] = pd.Series(pd.NA, index=master.index, dtype="Float64")
        # Distinguishes a test not yet run from a non-applicable test.
        master[f"{column.removesuffix('_score')}_status"] = "not_tested"

    # Keep duplicate source records visible rather than silently dropping epochs.
    epoch_key = ["Seed_id", "channel_start_time", "channel_end_time"]
    master["duplicate_epoch"] = master.duplicated(epoch_key, keep=False)
    master = master.sort_values(
        ["network", "station", "location", "channel", "channel_start_time"],
        na_position="last",
    ).reset_index(drop=True)

    # The SEED identifier already contains network, station, location and channel.
    master = master.drop(columns=["network", "station", "location", "channel"])
    status_columns = [
        f"{column.removesuffix('_score')}_status" for column in RESULT_COLUMNS
    ]
    protected = set(DOCUMENT_COLUMNS + status_columns + PROCESSING_METADATA_COLUMNS)

    # Drop globally empty fields; retain the explicitly requested future tests.
    master = master.replace(r"^\s*$", pd.NA, regex=True)
    empty_columns = [
        column
        for column in master.columns
        if column not in protected and master[column].isna().all()
    ]
    master = master.drop(columns=empty_columns)

    # Remove metadata columns that repeat another column for every record.
    kept = []
    for column in master.columns:
        if column not in protected and any(
            master[column].equals(master[previous]) for previous in kept
        ):
            continue
        kept.append(column)
    master = master[kept]
    other_columns = [
        column
        for column in master.columns
        if column not in DOCUMENT_COLUMNS + status_columns
    ]
    return master[DOCUMENT_COLUMNS + other_columns + status_columns]


def load_master_dataset(output_directory):
    """Load the CSV without interpreting identifiers, blanks or literal text."""
    path = Path(output_directory) / "master_inventory.csv"
    master = pd.read_csv(path, dtype=str, keep_default_na=False, low_memory=False)
    required = EPOCH_COLUMNS + ["inventory_file", "sample_rate_hz"]
    missing = set(required) - set(master.columns)
    if missing:
        raise ValueError(
            f"Master inventory lacks {sorted(missing)}; rebuild it with inventory.py"
        )
    return master


def update_processing_results(output_directory, updates):
    """Patch only PPSD fields in matching CSV/Excel rows, backing up both files.

    Each update contains the epoch key, inventory_file, and ppsd_* fields. Extra
    sheets, manual calibration results, annotations and existing styles survive.
    Both outputs are staged before replacing either original.
    """
    _update_master_results(output_directory, updates, prefix="ppsd_")


def update_pressure_test_results(output_directory, updates):
    """Patch pressure-comparison diagnostics and only the first test's result."""
    _update_master_results(
        output_directory,
        updates,
        prefix="pressure_expected_",
        allowed={TEST_COLUMNS[0], TEST_COLUMNS[0].removesuffix("_score") + "_status"},
    )


def update_infragravity_test_results(output_directory, updates):
    """Patch Godin diagnostics and only the predicted-infragravity test status."""
    column = TEST_COLUMNS[2]
    _update_master_results(
        output_directory,
        updates,
        prefix="infragravity_",
        allowed={column.removesuffix("_score") + "_status"},
    )


def update_nearby_test_results(output_directory, updates):
    """Patch peer diagnostics and only the second scientific check's status."""
    _update_master_results(
        output_directory,
        updates,
        prefix="nearby_",
        allowed={TEST_COLUMNS[1].removesuffix("_score") + "_status"},
    )


def update_paired_results(output_directory, updates):
    """Patch paired diagnostics only; preserve every scientific score/status."""
    _update_master_results(output_directory, updates, prefix="paired_")


def update_rayleigh_results(output_directory, updates):
    """Append provisional Rayleigh diagnostics; never change scientific scores."""
    _update_master_results(output_directory, updates, prefix="rayleigh_")


def update_local_p_results(output_directory, updates):
    """Patch uniquely keyed local_p_* diagnostics, never scores/check statuses.

    This explicit API does not run automatically at campaign/export time.
    Unlike legacy broad updaters it rejects duplicate full keys on either side.
    Existing CSV/Excel backup and transactional replacement semantics apply.
    """
    if not updates:
        return
    keys = EPOCH_COLUMNS + ["inventory_file"]
    seen = set()
    for update in updates:
        if not set(keys).issubset(update):
            raise ValueError("Local-P update requires all four epoch keys")
        if any(
            k not in keys and (not k.startswith("local_p_") or k.endswith("_score"))
            for k in update
        ):
            raise ValueError(
                "Local-P updates may change only local_p_* diagnostics, never scores"
            )
        key = tuple(str(update[k] or "") for k in keys)
        if not key[0] or not key[3] or key in seen:
            raise ValueError("Missing or duplicate local-P full epoch key")
        seen.add(key)
    with master_lock(output_directory):
        master = load_master_dataset(output_directory)
        for key in seen:
            mask = pd.Series(True, index=master.index)
            for k, v in zip(keys, key, strict=True):
                mask &= master[k].eq(v)
            if int(mask.sum()) != 1:
                raise ValueError("Local-P update needs exactly one matching master row")
        _patch_master(Path(output_directory), updates, "local_p_", ())


@contextmanager
def master_lock(output_directory):
    """Serialize master read-modify-write cycles between processes.

    Without it, two batches sharing one master (e.g. YL and YV) can each read
    the files, and the later commit silently drops the earlier one's fields.
    """
    output_directory = Path(output_directory)
    output_directory.mkdir(parents=True, exist_ok=True)
    with (output_directory / ".master_inventory.lock").open("a+") as lock:
        if not _lock_file(lock, wait=False):
            print("Waiting for another master update to finish", flush=True)
            _lock_file(lock, wait=True)
        try:
            yield
        finally:
            if fcntl is None:  # Windows does not promise to unlock on close
                lock.seek(0)
                msvcrt.locking(lock.fileno(), msvcrt.LK_UNLCK, 1)


def _lock_file(handle, *, wait):
    """Lock an open file exclusively; False when it is busy and not waiting."""
    if fcntl is not None:
        try:
            fcntl.flock(handle, fcntl.LOCK_EX | (0 if wait else fcntl.LOCK_NB))
        except BlockingIOError:
            return False
        return True
    # Windows: lock the first byte. LK_LOCK gives up after about 10 s; retry.
    handle.seek(0)
    while True:
        try:
            msvcrt.locking(
                handle.fileno(), msvcrt.LK_LOCK if wait else msvcrt.LK_NBLCK, 1
            )
            return True
        except OSError:
            if not wait:
                return False


def _update_master_results(output_directory, updates, *, prefix, allowed=()):
    """Stage and back up CSV/Excel updates restricted to the requested fields."""
    if not updates:
        return
    output_directory = Path(output_directory)
    with master_lock(output_directory):
        _patch_master(output_directory, updates, prefix, allowed)


def _patch_master(output_directory, updates, prefix, allowed):
    excel_path = output_directory / "master_inventory.xlsx"
    master = load_master_dataset(output_directory)
    keys = EPOCH_COLUMNS + ["inventory_file"]
    columns = list(
        dict.fromkeys(
            name
            for update in updates
            for name in update
            if name.startswith(prefix) or name in allowed
        )
    )
    if any(set(update) - set(keys + columns) for update in updates):
        raise ValueError(
            f"Updates may only change {prefix}* fields or {sorted(allowed)}"
        )
    for column in columns:
        if column not in master:
            master[column] = "not_processed" if column == "ppsd_status" else ""
    workbook = load_workbook(excel_path)
    worksheet = workbook["Master inventory"]
    headers = {cell.value: cell.column for cell in worksheet[1]}
    if not set(keys).issubset(headers):
        raise ValueError("Excel master is missing channel epoch/source keys")
    excel_rows = {}
    for row in worksheet.iter_rows(min_row=2):
        key = tuple(str(row[headers[name] - 1].value or "") for name in keys)
        excel_rows.setdefault(key, []).append(row[0].row)
    for column in columns:
        if column in headers:
            continue
        col = worksheet.max_column + 1
        headers[column] = col
        cell = worksheet.cell(1, col, column)
        cell._style = copy(worksheet.cell(1, 1)._style)
        worksheet.column_dimensions[cell.column_letter].width = min(
            max(len(column) + 2, 18), 40
        )
        if column == "ppsd_status":
            for row in range(2, worksheet.max_row + 1):
                worksheet.cell(row, col, "not_processed")
    for update in updates:
        key = tuple(str(update[name] or "") for name in keys)
        mask = pd.Series(True, index=master.index)
        for name, value in zip(keys, key, strict=True):
            mask &= master[name].eq(value)
        if not mask.any() or len(excel_rows.get(key, [])) != int(mask.sum()):
            raise ValueError(f"CSV/Excel epoch rows disagree for {key}")
        for column in columns:
            if column not in update:
                continue
            value = update[column]
            master.loc[mask, column] = "" if value is None else str(value)
            for row in excel_rows[key]:
                cell = worksheet.cell(row, headers[column])
                cell.value = value
                if cell.data_type == "f":
                    cell.data_type = "s"
                if column.endswith("_hz"):
                    cell.number_format = "0.000000"
                elif column.endswith("_segments") or column.endswith("_files"):
                    cell.number_format = "0"
        for column in columns:
            if column.endswith("_time") or column.endswith("_at"):
                letter = worksheet.cell(1, headers[column]).column_letter
                worksheet.column_dimensions[letter].width = max(
                    worksheet.column_dimensions[letter].width, 34
                )
    worksheet.auto_filter.ref = worksheet.dimensions
    _commit_master_files(output_directory, master, workbook)


def _commit_master_files(output_directory, master, workbook):
    """Stage and back up both files, restoring both on replacement failure."""
    csv_path = output_directory / "master_inventory.csv"
    excel_path = output_directory / "master_inventory.xlsx"
    timestamp = datetime.now(timezone.utc).strftime("%Y%m%dT%H%M%S%fZ")
    backup = output_directory / "backups" / timestamp
    with TemporaryDirectory(dir=output_directory) as staging:
        staged_csv = Path(staging) / csv_path.name
        staged_excel = Path(staging) / excel_path.name
        master.to_csv(staged_csv, index=False)
        workbook.save(staged_excel)
        workbook.close()
        backup.mkdir(parents=True)
        shutil.copy2(csv_path, backup / csv_path.name)
        shutil.copy2(excel_path, backup / excel_path.name)
        try:
            staged_csv.replace(csv_path)
            staged_excel.replace(excel_path)
        except OSError:
            shutil.copy2(backup / csv_path.name, csv_path)
            shutil.copy2(backup / excel_path.name, excel_path)
            raise
    print(f"Updated master CSV/Excel; previous files saved in {backup}", flush=True)


def synchronize_inventory(inventory_directory, output_directory=OUTPUT_DIRECTORY):
    """Append newly discovered channel/source epochs; preserve all existing cells.

    An existing metadata epoch is deliberately not refreshed automatically.
    This keeps processing results, manually entered scores and annotations
    intact when more network inventories are added to the input directory.
    """
    output_directory = Path(output_directory).expanduser().resolve()
    fresh = build_master_dataset(Path(inventory_directory).expanduser().resolve())
    with master_lock(output_directory):
        return _append_inventory(output_directory, fresh)


def _append_inventory(output_directory, fresh):
    csv_path = output_directory / "master_inventory.csv"
    excel_path = output_directory / "master_inventory.xlsx"
    if not csv_path.exists() and not excel_path.exists():
        _save_new_master(fresh, output_directory)
        return {"added_rows": len(fresh), "total_rows": len(fresh)}
    if not csv_path.exists() or not excel_path.exists():
        raise ValueError("Both existing master CSV and Excel must be present")
    master = load_master_dataset(output_directory)
    keys = EPOCH_COLUMNS + ["inventory_file"]
    known = set(map(tuple, master[keys].to_numpy()))
    fresh = fresh.astype(object).where(fresh.notna(), "")
    additions = fresh[
        [tuple(record) not in known for record in fresh[keys].to_numpy()]
    ].copy()
    if additions.empty:
        print("Inventory unchanged; preserved existing CSV/Excel", flush=True)
        return {"added_rows": 0, "total_rows": len(master)}
    book = load_workbook(excel_path)
    sheet = book["Master inventory"]
    headers = {cell.value: cell.column for cell in sheet[1]}
    excel_keys = [
        tuple(str(row[headers[name] - 1] or "") for name in keys)
        for row in sheet.iter_rows(min_row=2, values_only=True)
    ]
    if excel_keys != list(map(tuple, master[keys].to_numpy())):
        book.close()
        raise ValueError("CSV/Excel epoch rows disagree; inventory was not appended")
    # Use the existing long-text convention for newly appended records too.
    excel_additions = additions.copy()
    for name in additions:
        longest = (
            additions[name].map(lambda v: len(v) if isinstance(v, str) else 0).max()
        )
        if longest > 15000:
            for part in range((int(longest) + 14999) // 15000):
                target = name if part == 0 else f"{name}_part_{part + 1}"
                excel_additions[target] = additions[name].map(
                    lambda v, part=part: (
                        v[part * 15000 : (part + 1) * 15000]
                        if isinstance(v, str)
                        else v
                    )
                )
    for name in excel_additions:
        if name not in headers:
            col = sheet.max_column + 1
            headers[name] = col
            cell = sheet.cell(1, col, name)
            cell._style = copy(sheet.cell(1, 1)._style)
            sheet.column_dimensions[cell.column_letter].width = min(
                max(len(name) + 2, 18), 40
            )
    for _, record in excel_additions.iterrows():
        row_number = sheet.max_row + 1
        for name, value in record.items():
            cell = sheet.cell(row_number, headers[name])
            cell.value = None if value == "" else value
            if cell.data_type == "f":
                cell.data_type = "s"
    combined = pd.concat([master, additions], ignore_index=True).fillna("")
    sheet.auto_filter.ref = sheet.dimensions
    _commit_master_files(output_directory, combined, book)
    return {"added_rows": len(additions), "total_rows": len(combined)}


def save_master_dataset(master_df, output_directory):
    """Save matching CSV and Excel exports without modifying the input table.

    A fresh export never replaces an existing master: that would discard its
    processing results and manual scores without a backup.
    """
    output_directory = Path(output_directory)
    with master_lock(output_directory):
        return _save_new_master(master_df, output_directory)


def _save_new_master(master_df, output_directory):
    """Create a fresh export with master_lock already held by the caller."""
    master_csv = output_directory / "master_inventory.csv"
    master_excel = output_directory / "master_inventory.xlsx"
    existing = [str(path) for path in (master_csv, master_excel) if path.exists()]
    if existing:
        raise FileExistsError(
            f"Refusing to replace existing master {existing}; use "
            "`python -m compy_pressure.network inventory` to append new epochs, "
            "or choose a new output directory"
        )
    output_directory.mkdir(parents=True, exist_ok=True)
    master_df.to_csv(master_csv, index=False)
    print(f"Saved master dataset: {master_csv}")

    # Save an Excel copy; split long metadata to avoid Excel truncation.
    excel_df = master_df.copy()
    for column in master_df.columns:
        lengths = master_df[column].map(
            lambda value: len(value) if isinstance(value, str) else 0
        )
        if lengths.max() > 15000:
            for part in range((int(lengths.max()) + 14999) // 15000):
                name = column if part == 0 else f"{column}_part_{part + 1}"
                excel_df[name] = master_df[column].map(
                    lambda value, part=part: (
                        value[part * 15000 : (part + 1) * 15000]
                        if isinstance(value, str)
                        else value
                    )
                )

    with pd.ExcelWriter(master_excel, engine="openpyxl") as writer:
        excel_df.to_excel(writer, sheet_name="Master inventory", index=False)
        worksheet = writer.sheets["Master inventory"]
        worksheet.freeze_panes = "E2"
        worksheet.auto_filter.ref = worksheet.dimensions
        worksheet.row_dimensions[1].height = 75
        for cells in worksheet.iter_cols(min_row=1, max_row=1):
            cell = cells[0]
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            worksheet.column_dimensions[cell.column_letter].width = min(
                max(len(str(cell.value)) + 2, 18), 40
            )
        for row in worksheet.iter_rows(min_row=2):
            for cell in row:
                if cell.data_type == "f":
                    cell.data_type = "s"
    print(f"Saved Excel dataset: {master_excel}")
    return master_csv, master_excel


def main():
    """Build and export inventories from the command line."""
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument(
        "--inventory-dir",
        type=Path,
        default=XML_DIRECTORY,
        help="Folder containing StationXML files (searched recursively)",
    )
    parser.add_argument(
        "--output-dir",
        type=Path,
        default=OUTPUT_DIRECTORY,
        help="Folder for master_inventory.csv and master_inventory.xlsx",
    )
    args = parser.parse_args()
    master_df = build_master_dataset(args.inventory_dir)
    station_ids = master_df["Seed_id"].str.split(".").str[:2].str.join(".")
    print(
        f"Loaded {len(master_df)} channel epochs from {station_ids.nunique()} stations."
    )
    save_master_dataset(master_df, args.output_dir)


if __name__ == "__main__":
    main()
